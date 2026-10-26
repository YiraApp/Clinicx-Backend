"""Check that a functional test run may be reported, and write the report.

  python3 -I abdm_report.py plan  WORKBOOK --run ft|ui --entity government|private
  python3 -I abdm_report.py check WORKBOOK --run ft|ui --entity government|private [--meta meta.json]
  python3 -I abdm_report.py build WORKBOOK --run ft|ui --meta meta.json --evidence DIR --out REPORT.pdf
      [--manifest RUN.json] [--format pdf|docx|both] [--fonts DIR]
  python3 -I abdm_report.py tracker  WORKBOOK --run ft|ui [--entity government|private] [--plan INTEGRATION-PLAN.md] [--meta meta.json]
      [--set ID=building|built|tested|in-scope] [--screen CASE=ROUTE] [--build BUILD]
  python3 -I abdm_report.py findings WORKBOOK --run ft|ui --entity government|private --out REPORT.pdf

plan prints every case in sheet order with its group, its applicability for the
entity type and its status so far, then the totals: the run log is made from it.
check lists every mandatory case not PASS or CONDITIONAL PASS for the run type
and exits 1, or exits 0 when there is none. build refuses unless check passes, then writes the
report. A mandatory case NA passes only when meta.json's agreed_not_run names it.
--manifest is the run's manifest of case ids to REQUEST-IDs, printed per case.
build writes REPORT.pdf, the report to submit, and REPORT.docx beside it, the
editable copy; --format picks one of them. For a ui run it also writes
REPORT-screenshots.json, naming each figure's full size file for the annexure. The PDF uses Newsreader, Public Sans
and JetBrains Mono where it finds them (--fonts DIR, abdm-ft/fonts, then the
system's font folders), else Times-Roman, Helvetica and Courier.
tracker and findings read functionalities.json, which a module skill built per
functionality ships beside this script. tracker writes the functionality
tracker into INTEGRATION-PLAN.md, its state in tracker.json beside the
workbook. findings writes REPORT-findings.md: each FAIL and CONDITIONAL PASS,
its functionality, and whether a screen requirement in that functionality's
file cites it. build writes the findings too, when functionalities.json is there.
Needs openpyxl, Pillow, python-docx and reportlab
(pip install openpyxl pillow python-docx reportlab).
"""
import argparse, json, re, sys
from datetime import date
from pathlib import Path

import openpyxl

# The workbook layout shared.concept.test-workbook-write-back writes: the first
# column a name matches, with a "<column> (run)" column where NHA had text.
COLUMNS = {
    "ui": {"status": ["Status(PASS/FAIL)", "Status"], "response": ["Response of Request with Code"],
           "iterations": ["Iterations"], "remarks": ["Remarks", "Remark"]},
    "ft": {"status": ["FT Status"], "response": ["FT Response with Code"],
           "iterations": ["FT Iterations"], "remarks": ["FT Remarks"]},
}
# Header names vary by sheet and spacing: "Mandatory/ Optional" on HMIS and
# HPR, "Optional / Mandatory" on PHR; "Functionality", "Feature" or "Test Title".
TITLE = ["functionality", "feature", "test title"]
OPTIONAL = ("no", "optional", "non-mandatory", "non mandatory", "not mandatory")
FORBIDDEN = re.compile(r"pass(word|wd)|secret|token|credential|api[_ -]?key|private[_ -]?key", re.I)
# The statuses that pass a case, and how a multi step case's steps read in Remarks.
PASSING = ("PASS", "CONDITIONAL PASS")
STEPS = re.compile(r"\bSteps: (.+)$")
STEP = re.compile(r"^(\d+) (PASS|FAIL|needs tester|not run)(?: \((.*)\))?$")
IMAGES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff", ".heic", ".heif", ".pdf")
RUN_NAME = {"ui": "UI functional test", "ft": "API functional test"}
SECTIONS = ["Document control", "Executive summary", "ABDM details", "Results summary",
            "Scope and limitations", "Flow details", "Test environment and tools", "Test team",
            "Approach and methodology", "Test data", "Evidence", "Failures and iterations log",
            "Conclusion and sign-off", "Annexures"]
# The fixed fields of each tested case's evidence table, in order, under its id and title.
CASE_FIELDS = ["Result", "Expected result", "Observed", "Request IDs"]
# Where each case stands. Tested is passed, conditional pass or failed; applicable is
# tested, blocked or not tested; every case is applicable or not applicable.
TESTED = ("PASS", "CONDITIONAL PASS", "FAIL")
COUNTS = [("Cases in NHA's sheet", "Every case in the workbook."),
          ("Applicable to this integrator", "The cases less those not applicable."),
          ("Tested", "Run to a result: passed, conditional pass or failed."),
          ("Passed", "The Expected Result was seen."),
          ("Conditional pass", "The Expected Result was seen, with a gap beside it, listed with its condition."),
          ("Failed", "The Expected Result was not seen in three passes."),
          ("Blocked", "Started, then waiting for a tester, a device or test data."),
          ("Not tested", "Applicable and not run yet, NA on a mandatory case NHA has not agreed to, or an unknown status."),
          ("Not applicable", "Not for this entity type as the sheet's groups mark it, run or not; or NA agreed or optional.")]
# Screenshots print to be read and pack two rows to a page: a row is at most ROW_MM tall, the
# A4 frame (257 mm) less a case heading and fields (40 mm) and two captions and gaps (10 mm each),
# halved. A landscape one spans the text width up to that height, never under 140 mm wide; portrait
# ones sit three to a row at that height, never under 45 mm wide; a very tall one (over three times
# as tall as wide) prints 60 mm wide, cut into parts that each fit a page. Each is embedded with
# every pixel it has, up to what prints at 300 dots per inch.
ROW_MM = (257 - 40 - 2 * 10) / 2.0
SHOT_MM = {"wide_min_width": 140, "tall_min_width": 45, "very tall": 60, "part": 190}
DPI = 300
# A REQUEST-ID is a UUID. Only a UUID is printed as it is; anything else is masked.
UUID = re.compile(r"[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}")
# An error code or exception named in a response or remark.
ERROR = re.compile(r"\b(?:ABDM|MIS|HIS|HIU|HIP|CM|PHR|UHI)-\d{3,5}\b|\bHTTP [45]\d{2}\b|\b[A-Z]\w*(?:Exception|Error)\b")


def norm(text):
    return re.sub(r"\s+", " ", str(text if text is not None else "")).strip()


# Masking. This block is the same in abdm_workbook.py and abdm_report.py, so the
# sheet and the report hide the same things in the same way.
# A JWT: three base64url segments, the first a JSON header, so starting eyJ.
JWT = re.compile(r"\beyJ[\w-]+\.[\w-]+\.[\w-]*")
SECRET_WORDS = r"pass(?:word|wd)|secret|token|otp|credentials?|api[ _-]?key|private[ _-]?key"


def labelled(words):
    """A value after a label naming a secret: "client secret: x", "OTP 482913", "otp":"482913",
    "X-token: Bearer x". Only a value is taken, never the { or [ that opens an object or a list."""
    return re.compile(
        r"(\b[\w-]*(?:" + words + r")[\w-]*['\"]?\s*[:=]\s*['\"]?)(?![\[{])((?:Bearer\s+)?[^\s'\",;}]+)"
        r"|(\b(?:" + words + r")\b\s+(?:is\s+)?)((?=[^\s,;]*\d)[^\s,;]{4,})", re.I)


LABELLED = labelled(SECRET_WORDS)
BEARER = re.compile(r"\bBearer\s+(?!\[token\]|\[redacted\])[^\s'\",;}]+")
# A person's profile fields, by key, in JSON or in "key: value" and "key=value" text, in any case.
# A quoted value is taken whole; an unquoted one runs to the end of its clause: a comma, a
# semicolon, a full stop, a line end or a steps list.
PROFILE_KEYS = (r"firstName|middleName|lastName|fullName|name|dob|dayOfBirth|monthOfBirth|yearOfBirth|"
                r"address|pinCode|pincode|photo|profilePhoto|kycPhoto|email")
PROFILE = re.compile(
    r"(['\"]?\b(?i:" + PROFILE_KEYS + r")\b['\"]?\s*[:=]\s*)(?![\[{])"
    r"(\"(?:[^\"\\]|\\.)*\"|'[^']*'|[^\s,;}\]]+(?: (?!Steps:)[^\s,;.}\]]+)*)")
# A date of birth after its label, and any date written dd/mm/yyyy.
DATE = r"\d{1,2}[/-]\d{1,2}[/-]\d{4}|\d{4}-\d{2}-\d{2}"
BIRTH = re.compile(r"(\b(?i:dob|date of birth|birth ?date)\b['\"]?\s*[:=]?\s*['\"]?)(" + DATE + r")(?!\d)")
SLASHED = re.compile(r"(?<![\d/])\d{1,2}/\d{1,2}/\d{4}(?![\d/])")
ABHA_ADDRESS = re.compile(r"(?<![\w.-])([A-Za-z0-9._]+)@(sbx|abdm)\b(?!\.\w)", re.I)
# A digit run is not a number to mask when it is a segment of a UUID.
LEAD, TRAIL = r"(?<!\d)(?<![0-9a-fA-F]-)", r"(?!\d)(?!-[0-9a-fA-F])"
ABHA_NUMBER = re.compile(LEAD + r"\d{2}[ -]?\d{4}[ -]?\d{4}[ -]?(\d{4})" + TRAIL)
AADHAAR = re.compile(LEAD + r"\d{4}[ -]?\d{4}[ -]?(\d{4})" + TRAIL)
MOBILE = re.compile(r"(?<![\d+])(?<![0-9a-fA-F]-)(?:\+91[ -]?|91[ -]?|0)?[6-9](?:[ -]?\d){5}"
                    r"[ -]?(\d)[ -]?(\d)[ -]?(\d)[ -]?(\d)" + TRAIL)


def profile(m):
    quote = m.group(2)[0] if m.group(2)[:1] in ("'", '"') else ""
    return m.group(1) + quote + "[masked]" + quote


def address(m):
    """An ABHA address keeps the last four characters before the @, the rest as x. Four or fewer: all x."""
    name = m.group(1)
    if len(name) <= 4:
        return "x" * len(name) + "@" + m.group(2)
    return "x" * (len(name) - 4) + name[-4:] + "@" + m.group(2)


def mask(text):
    """A JWT to [token], a secret's value and a Bearer credential to [redacted], a profile field's
    value and a date of birth to [masked], a dd/mm/yyyy date to [date], an ABHA address to
    xxxxxx1234@sbx, ABHA numbers to XX-XXXX-XXXX-1234, any other 12 digit run to XXXX XXXX 1234,
    a mobile, with or without 0 or +91, to XXXXXX1234."""
    text = JWT.sub("[token]", str(text))
    text = LABELLED.sub(lambda m: (m.group(1) or m.group(3)) + "[redacted]", text)
    text = BEARER.sub("Bearer [redacted]", text)
    text = BIRTH.sub(lambda m: m.group(1) + "[masked]", text)
    text = PROFILE.sub(profile, text)
    text = SLASHED.sub("[date]", text)
    text = ABHA_ADDRESS.sub(address, text)
    text = ABHA_NUMBER.sub(r"XX-XXXX-XXXX-\1", text)
    text = AADHAAR.sub(r"XXXX XXXX \1", text)
    return MOBILE.sub(r"XXXXXX\1\2\3\4", text)
# End of the masking block.


# What meta.json may never hold as a value. An OTP is masked, not refused: it
# has expired by the time a report is written, and findings name the OTP step.
LABELLED_CREDENTIAL = labelled(r"pass(?:word|wd)|secret|token|credentials?|api[ _-]?key|private[ _-]?key")
LABEL = re.compile(r"\b(?:" + SECRET_WORDS + r")\b", re.I)


def load(path):
    return openpyxl.load_workbook(path, rich_text=True)


def header_row(ws):
    """The header row and, for each name, the first column carrying it."""
    for row in ws.iter_rows(min_row=1, max_row=15):
        names = [norm(c.value) for c in row]
        if "Test Case ID" in names:
            cols = {}
            for name, cell in zip(names, row):
                if name:
                    cols.setdefault(name, cell.column)
            return row[0].row, cols
    return None, {}


def cases(ws):
    """Case id to its first row. A repeated id gets #2, #3, as the writer does."""
    top, cols = header_row(ws)
    if top is None:
        return {}
    found, col = {}, cols["Test Case ID"]
    for r in range(top + 1, ws.max_row + 1):
        value = norm(ws.cell(r, col).value)
        if not value:
            continue
        base = key = value.replace(" ", "")
        n = 2
        while key in found:
            key, n = base + "#" + str(n), n + 1
        found[key] = r
    return found


def anchor(ws, row, col):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col)
    return ws.cell(row, col)


def index(wb):
    rows, twice = {}, []
    for ws in wb.worksheets:
        for case, row in cases(ws).items():
            if case in rows:
                twice.append(case)
            rows.setdefault(case, (ws, row))
    if twice:
        sys.exit("case id on more than one tab: " + ", ".join(twice))
    return rows


def find(cols, words):
    """The first column, in sheet order, whose name holds one of the words."""
    for word in words:
        for name, col in sorted(cols.items(), key=lambda x: x[1]):
            if word in name.lower() and not name.endswith(" (run)"):
                return col
    return None


def marking_col(cols):
    return next((col for name, col in sorted(cols.items(), key=lambda x: x[1])
                 if "mandatory" in name.lower() and "optional" in name.lower()), None)


def value(ws, row, top, cols, names, original):
    """The run's value. Where NHA's own cell had text, the writer put the run's value
    in "<column> (run)", and only that column is read for the row, even when blank."""
    col = next((cols[n] for n in names if n in cols), None)
    if col is None:
        return ""
    cell = anchor(ws, row, col)
    run_col = cols.get(norm(ws.cell(top, col).value) + " (run)")
    if run_col:
        if original is None:
            sys.exit("the workbook has a (run) column but no .before-abdm-run copy beside it, so NHA's "
                     "text cannot be told from the run's: restore the copy the writer made")
        nha = original[ws.title][cell.coordinate].value if ws.title in original.sheetnames else None
        if norm(nha):
            return norm(anchor(ws, row, run_col).value)
    return norm(cell.value)


def classify(text, entity):
    """(kind, label, note, pair ids) for one marking and entity type."""
    t = norm(text)
    low = t.lower()
    if not t:
        return "mandatory", "Not marked, treated as mandatory", "no marking: treated as mandatory", None
    if low.startswith("either"):
        ids = [i.replace(" ", "") for i in re.findall(r"\b[A-Z][A-Z0-9]*(?: ?_[A-Z0-9]+)+\b", t)]
        gov, priv = "gov" in low, "priv" in low
        if (entity == "government" and gov) or (entity == "private" and priv) or not (gov or priv):
            return "pair", "Either " + " or ".join(ids), "either rule: one of " + " or ".join(ids) + " must pass", ids
        return "optional", "Optional", "either rule names the other entity type: optional for " + entity, None
    if low in OPTIONAL or low.startswith(("optional", "non-mandatory", "non mandatory", "not mandatory", "no ")):
        return "optional", "Optional", None, None
    m = re.match(r"(mandatory|yes)\b[\s,:]*(.*)$", low)
    if not m:
        return "mandatory", "Unrecognised, treated as mandatory", "unrecognised marking: treated as mandatory", None
    rest = m.group(2)
    if not rest:
        return "mandatory", "Mandatory", None, None
    if rest.startswith("•"):
        # PHR's profile cases: "Mandatory" then a bullet list of the mandatory and optional fields.
        return "mandatory", "Mandatory", "lists mandatory and optional fields: the case is mandatory", None
    conditional = "(" in rest or re.search(r"\bif\b", rest) is not None
    who = re.match(r"for\s+(.*)$", rest)
    if who and m.group(1) == "mandatory":
        head = who.group(1).split("(")[0]
        gov, priv = "gov" in head, "priv" in head
        if (gov or priv) and not ((entity == "government" and gov) or (entity == "private" and priv)):
            return "optional", "Optional", "mandatory for the other entity type only: optional for " + entity, None
        if conditional:
            return "mandatory", "Mandatory (conditional)", "conditional: treated as mandatory", None
        return "mandatory", "Mandatory", "mandatory for " + entity, None
    if conditional:
        return "mandatory", "Mandatory (conditional)", "conditional: treated as mandatory", None
    return "mandatory", "Unrecognised, treated as mandatory", "unrecognised marking: treated as mandatory", None


def iterations(text):
    try:
        return int(float(text))
    except (TypeError, ValueError):
        return 0


def group_applies(ws, top, cols):
    """Each case row's group Applicable To, carried down from the row that heads the group."""
    acol, icol = cols.get("Applicable To"), cols.get("Test Case ID")
    if not acol or not icol:
        return {}
    out, current = {}, ""
    for r in range(top + 1, ws.max_row + 1):
        applies = norm(anchor(ws, r, acol).value)
        if applies:
            current = applies
        if norm(ws.cell(r, icol).value):
            out[r] = current
    return out


def group_names(ws, top, cols):
    """Each case row's group: the Function carried down from the row that heads the group,
    else the last section heading. A section heading is a row with no case id whose values sit in
    the first two columns only: a title, alone or after its number, as on the PHR sheet's
    "4.1 | Login with mobile number flow" and "| Edit Profile"."""
    fcol, icol = cols.get("Function"), cols.get("Test Case ID")
    if not icol:
        return {}
    first = min(cols.values())
    out, current = {}, ""
    for r in range(top + 1, ws.max_row + 1):
        name = norm(anchor(ws, r, fcol).value) if fcol else ""
        filled = [c for c in ws[r] if norm(c.value)]
        titles = [norm(c.value) for c in filled if not re.fullmatch(r"[\d.]+", norm(c.value))]
        if name:
            current = name
        elif (filled and not norm(ws.cell(r, icol).value) and all(c.column in (first, first + 1) for c in filled)
              and len(titles) == 1 and len(titles[0]) <= 120):
            current = titles[0]
        if norm(ws.cell(r, icol).value):
            out[r] = current or norm(ws.title)
    return out


def restricted_to(applies):
    """The one entity type a group is for, when it names only one: government or private."""
    low = applies.lower()
    gov, priv = "gov" in low, "priv" in low
    if "all" in low or "both" in low or gov == priv:
        return None
    return "government" if gov else "private"


def read_cases(path, run, entity):
    """Every case in sheet order, with its marking, the run's values and treatment."""
    wb = load(path)
    path = Path(path)
    # The copy the writer made before its first write: what NHA's own cells held.
    backup = path.with_name(path.stem + ".before-abdm-run" + path.suffix)
    original = load(backup) if backup.exists() else None
    out = []
    applies_by, groups_by = {}, {}
    for case, (ws, row) in index(wb).items():
        top, cols = header_row(ws)
        if ws.title not in applies_by:
            applies_by[ws.title] = group_applies(ws, top, cols)
            groups_by[ws.title] = group_names(ws, top, cols)
        mcol = marking_col(cols)
        marking = norm(anchor(ws, row, mcol).value) if mcol else ""
        tcol = find(cols, TITLE)
        title = norm(anchor(ws, row, tcol).value) if tcol else ""
        ecol = find(cols, ["expected result"])
        fcol = cols.get("Function")
        kind, label, note, pair = classify(marking, entity)
        rec = {"id": case, "sheet": ws.title, "title": title or case, "marking": marking,
               "kind": kind, "label": label, "note": note, "pair": pair,
               "expected": norm(anchor(ws, row, ecol).value) if ecol else "",
               "function": (norm(anchor(ws, row, fcol).value) if fcol else "") or ws.title}
        rec["applies"] = applies_by[ws.title].get(row, "")
        rec["group"] = groups_by[ws.title].get(row, "") or norm(ws.title)
        for key, names in COLUMNS[run].items():
            rec[key] = value(ws, row, top, cols, names, original)
        rec["raw_status"], rec["status"] = rec["status"], rec["status"].upper()
        rec["iterations"] = iterations(rec["iterations"])
        out.append(rec)
    # The either rule binds every case it names, whatever their own marking.
    groups = {}
    for rec in out:
        if rec["kind"] == "pair":
            groups[tuple(sorted(rec["pair"]))] = rec["pair"]
    by_id = {rec["id"]: rec for rec in out}
    for ids in groups.values():
        for i in ids:
            if i in by_id:
                other = [x for x in ids if x != i]
                by_id[i].update(kind="pair", pair=ids, label="Either, with " + " or ".join(other),
                                note="either rule: one of " + " or ".join(ids) + " must pass")
    # A group the sheet marks for one entity type only, such as M1's demographic authentication
    # for Government, is not for the other one, whatever its cases' own markings say.
    for rec in out:
        only = restricted_to(rec["applies"])
        if only and only != entity:
            rec.update(kind="optional", pair=None, label="Not for a " + entity + " entity", na=True,
                       note="the group applies to " + rec["applies"] + " only: not for a " + entity + " entity")
    groups = {k: v for k, v in groups.items() if any(by_id[i]["kind"] == "pair" for i in v if i in by_id)}
    return out, list(groups.values())


def passes_gate(r, agreed):
    """One case through the gate: PASS, a CONDITIONAL PASS with its condition written down, or NA on a
    case NHA agreed is not run. The gate and the tracker both read this."""
    if r["status"] == "CONDITIONAL PASS":
        return bool(STEPS.sub("", r["remarks"]).strip())
    return r["status"] in PASSING or (r["status"] == "NA" and r["id"] in agreed)


def gate(recs, groups, agreed):
    """Mandatory cases not PASS or CONDITIONAL PASS. NA passes only on a case NHA agreed not to run."""
    blocking = []
    by_id = {r["id"]: r for r in recs}

    def through(r):
        return passes_gate(r, agreed)

    def shown_status(r):
        if r["status"] == "CONDITIONAL PASS" and not through(r):
            return "CONDITIONAL PASS with no condition in Remarks"
        return r["status"] or "blank"

    for r in recs:
        if r["kind"] == "mandatory" and not through(r):
            blocking.append(r["id"] + ": " + shown_status(r))
    for ids in groups:
        present = [by_id[i] for i in ids if i in by_id]
        if present and not any(through(r) for r in present):
            blocking.append(" or ".join(r["id"] for r in present) + ": " +
                            ", ".join(shown_status(r) for r in present))
    return blocking


def agreed_cases(meta):
    """meta.json's agreed_not_run: case id to who at NHA agreed it is not run, and when."""
    agreed = meta.get("agreed_not_run") or {}
    if not isinstance(agreed, dict) or not all(isinstance(v, str) and v.strip() for v in agreed.values()):
        sys.exit('agreed_not_run must map each case id to who at NHA agreed and when, such as '
                 '{"CRT_ABHA_101": "<NAME>, NHA, email of 2026-10-21"}')
    return {norm(k).replace(" ", ""): norm(v) for k, v in agreed.items()}


def run_check(path, run, entity, agreed):
    recs, groups = read_cases(path, run, entity)
    if not recs:
        sys.exit("no Test Case ID column with ids in " + Path(path).name + ": nothing to report on")
    for r in recs:
        if r["note"]:
            print("marking " + r["id"] + ' "' + r["marking"] + '": ' + r["note"])
        r["agreed"] = agreed.get(r["id"]) if r["status"] == "NA" and r["kind"] in ("mandatory", "pair") else None
        if r["agreed"]:
            print("agreed not run " + r["id"] + ": " + mask(r["agreed"]))
    blocking = gate(recs, groups, agreed)
    for line in blocking:
        print("not passed " + line)
    mandatory = sum(1 for r in recs if r["kind"] == "mandatory") + len(groups)
    print(str(len(blocking)) + " of " + str(mandatory) + " mandatory case(s) for a " + entity +
          " entity not PASS in the " + run + " run")
    for r in recs:
        if r["status"] == "CONDITIONAL PASS":
            print("conditional pass " + r["id"] + ": " + mask(r["remarks"] or "no condition recorded"))
    return recs, groups, blocking


def applicability(r):
    """What a case is for this entity type, in the words plan prints."""
    if r.get("na"):
        return "not applicable (group " + r["group"] + ", Applicable To: " + r["applies"] + ")"
    if r["kind"] == "pair":
        what = "mandatory, either " + " or ".join(r["pair"])
    elif r["kind"] == "mandatory":
        what = "conditional" if "conditional" in r["label"].lower() else "mandatory"
    else:
        what = "optional"
    low = r["applies"].lower()
    if r["applies"] and not restricted_to(r["applies"]) and not ("all" in low or "both" in low or
                                                                  ("gov" in low and "priv" in low)):
        # An Applicable To that names neither entity type, such as "PHR App", is read as all.
        what += "; applies: " + r["applies"] + " (read as all)"
    return what


def settle(recs, groups):
    """Mark what the counts need: a mandatory NA without NHA's agreement is still open, and an
    either pair met by one case leaves the other's result beside it, not a mandatory failure.
    Returns the mandatory cases to count: each mandatory case, and one case for each either pair."""
    by_id = {r["id"]: r for r in recs}

    def through(r):
        return r["status"] in PASSING or bool(r.get("agreed"))

    mandatory = [r for r in recs if r["kind"] == "mandatory"]
    for ids in groups:
        members = [by_id[i] for i in ids if i in by_id and by_id[i]["kind"] == "pair"]
        if not members:
            continue
        met = [r for r in members if through(r)]
        for r in members:
            r["either"] = " or ".join(x["id"] for x in members)
            r["pair_met"] = bool(met) and r not in met
        tried = [r for r in members if r["status"] in TESTED or r["remarks"].lower().startswith("needs tester")]
        mandatory.append((met or tried or members)[0])
    for r in recs:
        r["open_na"] = (r["status"] == "NA" and r["kind"] in ("mandatory", "pair") and not r.get("agreed")
                        and not r.get("pair_met"))
    return sorted(mandatory, key=recs.index)


def standing(r):
    """Where a case stands: passed, conditional pass, failed, blocked, not tested or not applicable.
    What plan classifies not applicable stays so, run or not; a mandatory NA nobody agreed is not tested."""
    status = r["status"]
    if r.get("na"):
        return "not applicable"
    if status in TESTED:
        return {"PASS": "passed", "CONDITIONAL PASS": "conditional pass", "FAIL": "failed"}[status]
    if status == "NA":
        return "not tested" if r.get("open_na") else "not applicable"
    if not status and r["remarks"].lower().startswith("needs tester"):
        return "blocked"
    return "not tested"


def tally(recs):
    """The summary counts, checked to add up: applicable is tested, blocked and not tested;
    tested is passed, conditional pass and failed; the sheet is applicable and not applicable."""
    c = {"cases": len(recs)}
    for name in ("passed", "conditional pass", "failed", "blocked", "not tested", "not applicable"):
        c[name] = sum(1 for r in recs if standing(r) == name)
    c["tested"] = sum(1 for r in recs if r["status"] in TESTED and not r.get("na"))
    c["applicable"] = sum(1 for r in recs if not r.get("na") and not (r["status"] == "NA" and not r.get("open_na")))
    if not (c["applicable"] == c["tested"] + c["blocked"] + c["not tested"]
            and c["tested"] == c["passed"] + c["conditional pass"] + c["failed"]
            and c["cases"] == c["applicable"] + c["not applicable"]):
        sys.exit("the summary counts do not add up: " + json.dumps(c))
    return c


def run_plan(path, run, entity):
    """Every case in sheet order: id, group, applicability, status so far. Then the totals."""
    recs, _ = read_cases(path, run, entity)
    if not recs:
        sys.exit("no Test Case ID column with ids in " + Path(path).name + ": nothing to plan")
    counts, statuses = {}, {}
    for r in recs:
        what = applicability(r)
        status = r["status"] or "open"
        print(" | ".join([r["id"], r["group"], what, status]))
        key = what.split(" (")[0].split(",")[0].split(";")[0]
        counts[key] = counts.get(key, 0) + 1
        statuses[status] = statuses.get(status, 0) + 1
    print("total " + str(len(recs)) + " cases: " + ", ".join(str(v) + " " + k for k, v in counts.items()) +
          "; status: " + ", ".join(str(v) + " " + k for k, v in statuses.items()))


def forbidden(meta, where="meta.json"):
    found = []
    if isinstance(meta, dict):
        for k, v in meta.items():
            if FORBIDDEN.search(str(k)):
                found.append(where + "." + str(k))
            found += forbidden(v, where + "." + str(k))
    elif isinstance(meta, list):
        for n, v in enumerate(meta):
            found += forbidden(v, where + "[" + str(n) + "]")
    elif isinstance(meta, str):
        # A value holding a token, or a credential's label with its value, such as "client secret: x".
        if JWT.search(meta) or LABELLED_CREDENTIAL.search(meta) or (
                re.search(r"\bclient[ _-]?secret\b", meta, re.I)):
            found.append(where + " (its value)")
    return found


def request_ids(path):
    """The run's manifest: case id to its REQUEST-IDs. Takes {"<id>": [..]}, {"<id>": "<rid>"},
    {"<id>": {"request_ids": [..]}}, the same under "cases", or a list of {"id": .., "request_ids": [..]}."""
    if not path:
        return None
    data = json.loads(Path(path).read_text())
    bad = forbidden(data, "manifest")
    if bad:
        sys.exit("the manifest carries a credential, which never goes in a report: " + ", ".join(bad))
    if isinstance(data, dict) and isinstance(data.get("cases"), (dict, list)):
        data = data["cases"]
    if isinstance(data, list):
        data = {str(x.get("id") or x.get("case") or ""): x for x in data if isinstance(x, dict)}
    out = {}
    for case, ids in (data.items() if isinstance(data, dict) else []):
        if isinstance(ids, dict):
            ids = ids.get("request_ids") or ids.get("requestIds") or ids.get("request_id") or []
        if isinstance(ids, str):
            ids = [ids]
        # A REQUEST-ID is a UUID; anything else in the field is not printed.
        out[norm(case).replace(" ", "")] = [norm(i) if UUID.fullmatch(norm(i)) else "[invalid id]"
                                            for i in ids if norm(i)]
    return out


def days(start, end):
    try:
        return (date.fromisoformat(end) - date.fromisoformat(start)).days + 1
    except (TypeError, ValueError):
        return ""




def codes(*texts):
    """The error codes and exceptions named in a response or remark."""
    found = []
    for text in texts:
        for c in ERROR.findall(text or ""):
            if c not in found:
                found.append(c)
    return ", ".join(found)


def shown(row, header=False):
    """A table row as text. A value beside a label naming a secret, such as "OTP used", is never shown."""
    secret = len(row) == 2 and not header and LABEL.search(str(row[0] or ""))
    return ["[redacted]" if secret and c == 1 and text else str(text if text is not None else "")
            for c, text in enumerate(row)]


def step_rows(remarks):
    """A multi step case's steps, read back from the "Steps: ..." the writer put at the end of Remarks."""
    found = STEPS.search(remarks or "")
    rows = []
    for part in (found.group(1).split("; ") if found else []):
        m = STEP.match(part.strip())
        if m:
            rows.append([m.group(1), m.group(2), m.group(3) or ""])
    return rows


def shot_rows(shots):
    """Screenshots into rows, in name order within each shape: a landscape one alone across the
    text width, portrait ones three to a row, a very tall one alone. Shapes never share a row."""
    kinds = []
    for shot in shots:
        if shot["kind"] not in kinds:
            kinds.append(shot["kind"])
    rows = []
    for kind in kinds:
        these = [s for s in shots if s["kind"] == kind]
        per = 3 if kind == "tall" else 1
        rows += [these[i:i + per] for i in range(0, len(these), per)]
    return rows


def unreadable(path):
    """Why a screenshot cannot be placed, or None when Pillow opens it."""
    from PIL import Image
    try:
        with Image.open(path) as im:
            im.load()
            if not im.size[0] or not im.size[1]:
                return "no pixels"
        return None
    except Exception as e:
        return type(e).__name__


def content(path, run, meta, evidence, recs, groups, manifest=None, draft=None, out=None):
    """What the report says, as a list of blocks. The DOCX and the PDF draw the same blocks."""
    from PIL import Image

    dm = meta.get("document", {})
    dates = meta.get("test_dates", {})
    milestone = meta.get("milestone", "")
    entity = meta["entity_type"]
    label = RUN_NAME[run]
    title = milestone + " " + label + " report"
    if draft:
        # A draft carries the mark on every page, so it is never mistaken for the report to submit.
        title = "Draft " + title
    rids = request_ids(manifest)
    safe = set()  # text printed unmasked: a list of REQUEST-IDs that are all UUIDs
    blocks = []

    def add(*block):
        blocks.append(block)

    def table(rows, header=None):
        add("table", rows, header)

    counted = settle(recs, groups)
    everything = tally(recs)
    mandatory = tally(counted)
    conditional = [r for r in recs if r["status"] == "CONDITIONAL PASS"]
    agreed = [r for r in recs if r.get("agreed")]
    tested = [r for r in recs if standing(r) in ("passed", "conditional pass", "failed")]
    untested = [r for r in recs if r not in tested]
    log = [r for r in recs if standing(r) in ("failed", "blocked", "conditional pass")]
    findings = meta.get("findings", {})
    ready = ("Ready for sandbox exit." if not draft else
             str(len(draft)) + " mandatory case" + (" is" if len(draft) == 1 else "s are") + " still open.")
    functions = []
    for r in recs:
        if r["group"] not in functions:
            functions.append(r["group"])
    by_function = [r for fn in functions for r in recs if r["group"] == fn]

    def condition(r):
        return STEPS.sub("", r["remarks"]).strip() or "Not recorded"

    def needed(r):
        """For a case not tested: what it waits for, or why it does not apply."""
        if r.get("agreed"):
            return "Agreed not run: " + r["agreed"]
        if r.get("na"):
            ran = ("Run, result " + r["status"] + ". ") if r["status"] in TESTED else ""
            return ran + "The group applies to " + r["applies"] + " only."
        if r["status"] and r["status"] not in TESTED + ("NA",):
            return "Unrecognised status: " + r["raw_status"]
        text = STEPS.sub("", r["remarks"]).strip()
        text = re.sub(r"^needs tester:\s*", "", text, flags=re.I)
        if r.get("open_na"):
            return ("Recorded NA without NHA's agreement: run it, or add the agreement to agreed_not_run."
                    + (" Remarks: " + text if text else ""))
        if not text and standing(r) == "blocked":
            text = "; ".join(s[2] for s in step_rows(r["remarks"]) if s[1] == "needs tester" and s[2])
        return text or ("Not run yet." if standing(r) == "not tested" else "Not recorded.")

    def shown_status(r):
        if r.get("na") and r["status"] in TESTED:
            return "Run, not applicable"
        return standing(r).capitalize()

    def either(r):
        return (" Either rule: " + r["either"] + " satisfies it.") if r.get("pair_met") else ""

    def passes(r):
        n = r["iterations"] or 1
        return str(n) + (" pass" if n == 1 else " passes")

    def rid_text(r):
        if rids is None:
            return "No manifest given"
        ids = rids.get(r["id"], [])
        if not ids:
            return "None in the manifest"
        text = ", ".join(ids)
        if all(UUID.fullmatch(i) for i in ids):
            safe.add(text)
        return text

    # 1. Document control
    add("h1", SECTIONS[0])
    table([["Document title", title], ["Document id", dm.get("id", "FT-" + milestone)],
           ["Version", dm.get("version", "")], ["Prepared by", dm.get("prepared_by", "")],
           ["Reviewed by", dm.get("reviewed_by", "")], ["Approved by", dm.get("approved_by", "")],
           ["Released by", dm.get("released_by", "")], ["Release date", dm.get("release_date", "")]])
    add("h2", "Change history")
    history = dm.get("change_history", [])
    table([[h.get("version", ""), h.get("date", ""), h.get("by", ""), h.get("change", "")] for h in history]
          or [["", "", "", ""]], header=["Version", "Date", "By", "Change"])
    add("h2", "Distribution")
    dist = dm.get("distribution", [])
    table([[d.get("name", ""), d.get("organisation", ""), d.get("role", "")] if isinstance(d, dict)
           else [str(d), "", ""] for d in dist] or [["", "", ""]], header=["Name", "Organisation", "Role"])

    # 2. Executive summary: the counts, each term in one line, then whether the run is done.
    add("h1", SECTIONS[1])
    keys = ["cases", "applicable", "tested", "passed", "conditional pass", "failed", "blocked", "not tested",
            "not applicable"]
    add("counts", [[name, everything[k], mandatory[k]] for (name, _), k in zip(COUNTS, keys)],
        ["Measure", "All cases", "Mandatory"],
        [[name, meaning.replace("not applicable", "not applicable to a " + entity + " entity")
          if name == "Applicable to this integrator" else meaning] for name, meaning in COUNTS], ready)

    # 3. ABDM details
    add("h1", SECTIONS[2])
    recordings = meta.get("recordings", [])
    table([["Entity name", meta.get("organisation", "")], ["Product name", meta.get("product", "")],
           ["Category", meta.get("category", "")], ["Role", meta.get("role", "")], ["Entity type", entity],
           ["Milestone", milestone], ["API version", meta.get("api_version", "V3")],
           ["Run type", run + ", " + label],
           ["Test execution dates", dates.get("start", "") + " to " + dates.get("end", "")],
           ["Number of days for testing", meta.get("test_days", days(dates.get("start"), dates.get("end")))],
           ["Most passes any case took", max([r["iterations"] for r in recs] or [0])],
           ["Test executed by", meta.get("executed_by", "")],
           ["Recordings", str(len(recordings)) + " listed in the annexures" if recordings else "None"]])

    # 4. Results summary: by function, the conditional passes, the agreed NA, then every case not tested.
    add("h1", SECTIONS[3])
    add("h2", "Results by function")
    rows = []
    for fn in functions:
        c = tally([r for r in recs if r["group"] == fn])
        rows.append([fn, c["cases"], c["passed"], c["conditional pass"], c["failed"], c["blocked"],
                     c["not tested"], c["not applicable"]])
    table(rows, header=["Function", "Cases", "Passed", "Conditional pass", "Failed", "Blocked", "Not tested",
                        "Not applicable"])
    if conditional:
        add("h2", "Conditional passes")
        table([[r["id"], r["title"], condition(r)] for r in conditional], header=["Test case ID", "Title", "Condition"])
    if agreed:
        add("h2", "Mandatory cases NHA agreed not to run")
        table([[r["id"], r["title"], r["agreed"]] for r in agreed], header=["Test case ID", "Title", "Agreed by and when"])
    if untested:
        add("h2", "Cases not tested")
        table([[r["id"], r["title"], r["label"], shown_status(r), needed(r)]
               for r in by_function if r in untested],
              header=["Test case ID", "Title", "Marking", "Status", "Needed, or why not applicable"])

    # 5. Scope and limitations
    add("h1", SECTIONS[4])
    apps = meta.get("applications", [])
    table([[a.get("name", ""), a.get("url", "")] if isinstance(a, dict) else [str(a), ""] for a in apps]
          or [["", ""]], header=["Application", "URL"])
    sheets = sorted(set(r["sheet"] for r in recs), key=[r["sheet"] for r in recs].index)
    table([["Workbook", Path(path).name], ["Tabs", ", ".join(sheets)],
           ["Limitations", " ".join(["The results hold for the build and configuration tested from " +
                                     dates.get("start", "") + " to " + dates.get("end", "") + "."] +
                                    [str(x) for x in meta.get("limitations", [])])]])

    # 6. Flow details
    add("h1", SECTIONS[5])
    flows = meta.get("flows", [])
    if flows:
        table([[f.get("name", ""), ", ".join(f.get("abha_addresses", [])), ", ".join(f.get("hi_types", []))]
               for f in flows if isinstance(f, dict)],
              header=["Flow", "ABHA address used", "HI types created or tested"])
    else:
        add("para", "No linking or consent flow in this milestone.")

    # 7. Test environment and tools
    add("h1", SECTIONS[6])
    env = dict(meta.get("environment", {}))
    names = [("browser", "Browser"), ("browser_version", "Browser version"), ("playwright", "Playwright"),
             ("phr_app", "PHR app"), ("phr_app_version", "PHR app version"), ("emulator", "Emulator"),
             ("model", "Model that ran the agent")]
    rows = [[name, env.pop(key, "Not used")] for key, name in names]
    rows += [[key, val] for key, val in env.items()]
    table(rows, header=["Tool", "Version or value"])

    # 8. Test team
    add("h1", SECTIONS[7])
    team = meta.get("team", [])
    table([[m.get("name", ""), m.get("role", ""), m.get("qualifications", "")] for m in team if isinstance(m, dict)]
          or [[meta.get("executed_by", ""), "Tester", ""]], header=["Name", "Role", "Qualifications"])

    # 9. Approach and methodology: one line each.
    add("h1", SECTIONS[8])
    table([["Runner", "An agent ran each case with the module's test skill and wrote the result into the workbook."],
           ["Passes", "At most three per case. Iterations is the number of passes a case took."],
           ["Approvals", "A person gave every OTP, consent and approval as the case needed it. The agent kept no OTP."],
           ["PHR app", "A case through the PHR app is proved by one PHR app screenshot of the consent or the record."],
           ["Run type", run + ": Playwright drove the screens, with a screenshot per step." if run == "ui" else
            run + ": the calls to the sandbox and the callbacks decide each case. The response is the evidence."],
           ["Request IDs", "A fresh one on every call. Look each up in the gateway's own record."]])

    # 10. Test data
    add("h1", SECTIONS[9])
    data = meta.get("test_data", [])
    if isinstance(data, dict):
        data = [{"item": k, "value": v} for k, v in data.items()]
    table([[d.get("item", ""), d.get("value", "")] for d in data] or [["", ""]], header=["Item", "Value"])

    # 11. Evidence: tested cases only, by function. Each: a result line, the expected and observed
    # results, its REQUEST-IDs, any steps, the response for an ft run, then its screenshots in rows.
    add("h1", SECTIONS[10])
    figures = []
    if not tested:
        add("para", "No case was tested.")
    for fn in functions:
        these = [r for r in recs if r["group"] == fn and r in tested]
        if not these:
            continue
        add("group", fn)
        for r in these:
            observed = findings.get(r["id"]) or STEPS.sub("", r["remarks"]).strip() or "None recorded."
            result = r["status"] + ", " + passes(r) + ". Marking: " + r["label"] + "." + either(r)
            if r["status"] == "CONDITIONAL PASS" and condition(r) != observed:
                result += " Condition: " + condition(r) + ("" if condition(r).endswith(".") else ".")
            err = codes(r["response"], r["remarks"]) if r["status"] != "PASS" or r["iterations"] > 1 else ""
            if err:
                result += " Error code: " + err + "."
            values = [result, r["expected"] or "Not given in the workbook", observed, rid_text(r)]
            shots, unread = [], []
            if run == "ui":
                folder = Path(evidence) / r["id"]
                files = sorted(f for f in folder.iterdir() if f.suffix.lower() in IMAGES) if folder.is_dir() else []
                for step, shot in enumerate(files, 1):
                    if unreadable(shot):
                        # A file Pillow cannot open is named, never dropped and never a crash.
                        unread.append("Screenshot could not be read: " + shot.name + " (" + unreadable(shot) + ")")
                        continue
                    with Image.open(shot) as im:
                        w, h = im.size
                    shots.append({"path": shot, "step": step, "w": w, "h": h,
                                  "kind": "very tall" if h > 3 * w else "tall" if h > 1.2 * w else "wide"})
            rows = shot_rows(shots)
            for shot in [s for row in rows for s in row]:
                # Figures are numbered in the order they print.
                figures.append({"figure": len(figures) + 1, "case": r["id"], "step": shot["step"],
                                "file": str(shot["path"])})
                shot["figure"] = len(figures)
            add("case", r, values, r["response"] if run == "ft" else "", rows, step_rows(r["remarks"]), unread)

    # 12. Failures and iterations log: failed, blocked and conditional pass only.
    add("h1", SECTIONS[11])
    if log:
        table([[r["id"], standing(r).capitalize(), r["iterations"] or "",
                (condition(r) if r["status"] == "CONDITIONAL PASS" else needed(r) if standing(r) == "blocked"
                 else STEPS.sub("", r["remarks"]).strip() or r["response"] or "Not recorded") + either(r),
                codes(r["response"], r["remarks"]) or "None found"] for r in log],
              header=["Test case ID", "Status", "Iterations", "Failure reason", "Error code"])
    else:
        add("para", "No case failed, is blocked or passed with a condition.")

    # 13. Conclusion and sign-off
    add("h1", SECTIONS[12])
    add("ready", ready)
    table([["Prepared by", dm.get("prepared_by", ""), "", ""], ["Reviewed by", dm.get("reviewed_by", ""), "", ""],
           ["Approved by", dm.get("approved_by", ""), "", ""]], header=["Sign-off", "Name", "Date", "Signature"])

    # 14. Annexures
    add("h1", SECTIONS[13])
    table([[x.get("title", ""), x.get("link", "")] if isinstance(x, dict) else [str(x), ""] for x in recordings]
          or [["", ""]], header=["Recording", "Link"])
    listing = Path(out).stem + "-screenshots.json" if out else "Not written"
    table([["Completed workbook", Path(path).name],
           ["REQUEST-ID manifest", Path(manifest).name if manifest else "Not given"],
           ["Full size screenshots", (str(evidence) + ", each figure listed in " + listing if run == "ui"
                                      else "Not used: the responses are the evidence")]],
          header=["Annexure", "File"])

    return {
        "title": title, "label": label, "run": run, "milestone": milestone, "entity": entity, "safe": safe,
        "draft": draft or [], "figures": figures, "counts": everything, "mandatory": mandatory,
        "product": meta.get("product", ""), "organisation": meta.get("organisation", ""), "blocks": blocks,
        "cover": [["Product", meta.get("product", "")], ["Organisation", meta.get("organisation", "")],
                  ["Milestone", milestone], ["Report type", label + " (run type " + run + ")"],
                  ["Document id", dm.get("id", "FT-" + milestone)], ["Version", dm.get("version", "")],
                  ["Release date", dm.get("release_date", "")]],
        "facts": [["Document id", dm.get("id", "FT-" + milestone)], ["Version", dm.get("version", "")],
                  ["Release date", dm.get("release_date", "")], ["Run type", run + ", " + label],
                  ["Entity type", entity], ["Test dates", dates.get("start", "") + " to " + dates.get("end", "")],
                  ["Executed by", meta.get("executed_by", "")]],
        "confidential": ("Confidential. Aadhaar, ABHA and mobile numbers are masked, and the report "
                         "holds no credential, token or OTP."),
    }


def shot_parts(shot, text_mm, gap_mm=4):
    """How a screenshot prints, as (crop box or None, width mm, height mm) for each piece: a
    landscape one across the text width, no taller than a row and never under 140 mm wide; a
    portrait one a row tall, never under 45 mm wide and never over a third of the row; a very tall
    one 60 mm wide, cut into pieces that fit a page."""
    w, h = float(shot["w"]), float(shot["h"])
    if shot["kind"] == "wide":
        width = min(text_mm, max(SHOT_MM["wide_min_width"], ROW_MM * w / h))
        return [(None, width, width * h / w)]
    if shot["kind"] == "tall":
        slot = (text_mm - 2 * gap_mm) / 3.0
        width = min(slot, max(SHOT_MM["tall_min_width"], ROW_MM * w / h))
        return [(None, width, width * h / w)]
    width = SHOT_MM["very tall"]
    pieces = max(1, int(-(-(width * h / w) // SHOT_MM["part"])))
    step = int(-(-h // pieces))
    out = []
    for i in range(pieces):
        top, bottom = i * step, min(int(h), (i + 1) * step)
        out.append(((0, top, int(w), bottom), width, width * (bottom - top) / w))
    return out


def embedded(shot, crop, w_mm, h_mm):
    """The screenshot as it goes in the file: every source pixel, unless that is more than prints
    at 300 dots per inch, when it is scaled to that. PNG stays lossless; JPEG keeps quality 95.
    The full size file stays in the evidence folder for the annexure bundle."""
    from io import BytesIO
    from PIL import Image
    cap = (max(1, int(w_mm / 25.4 * DPI)), max(1, int(h_mm / 25.4 * DPI)))
    with Image.open(shot["path"]) as im:
        jpeg = im.format == "JPEG"
        if crop is None and im.format in ("PNG", "JPEG") and im.size[0] <= cap[0]:
            return BytesIO(Path(shot["path"]).read_bytes())
        im.load()
        if crop is not None:
            im = im.crop(crop)
        if im.mode not in ("RGB", "RGBA", "L"):
            im = im.convert("RGBA" if "A" in im.mode or im.mode == "P" else "RGB")
        if im.size[0] > cap[0]:
            im = im.resize((cap[0], max(1, int(im.size[1] * cap[0] / float(im.size[0])))), Image.LANCZOS)
        buf = BytesIO()
        if jpeg:
            im.convert("RGB").save(buf, "JPEG", quality=95)
        else:
            im.save(buf, "PNG", optimize=True)
    buf.seek(0)
    return buf


def caption(shot, piece, pieces, short=True):
    text = ("Fig " if short else "Figure ") + str(shot["figure"]) + ", step " + str(shot["step"])
    return text + (", part " + str(piece) + " of " + str(pieces) if pieces > 1 else "")


def pieces_of(text, size=600):
    """Long text in pieces of about size characters, broken at a space, so each fits a page."""
    text = str(text if text is not None else "")
    out = []
    while len(text) > size:
        cut = text.rfind(" ", 0, size)
        cut = cut if cut > size // 2 else size
        out.append(text[:cut])
        text = text[cut:].lstrip()
    return out + [text]


def write_docx(m, out):
    """The editable copy, in Word's format: the same blocks, with Word's contents and page fields."""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin

    def field(paragraph, code, placeholder=""):
        """A Word field, such as PAGE or a TOC, that Word fills in when the document opens."""
        def char(kind):
            r = paragraph.add_run()
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        char("begin")
        r = paragraph.add_run()
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " " + code + " "
        r._r.append(instr)
        char("separate")
        paragraph.add_run(placeholder)
        char("end")

    def table(rows, header=None):
        t = doc.add_table(rows=0, cols=len(header or rows[0]))
        t.style = "Table Grid"
        for n, row in enumerate(([header] if header else []) + rows):
            cells = t.add_row().cells
            for c, (cell, text) in enumerate(zip(cells, shown(row, header and n == 0))):
                cell.text = text
                if (header and n == 0) or (len(row) == 2 and c == 0 and not header):
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.bold = True
            if header and n == 0:
                # The header row repeats on every page the table runs to.
                tr = t.rows[0]._tr.get_or_add_trPr()
                el = OxmlElement("w:tblHeader")
                el.set(qn("w:val"), "true")
                tr.append(el)
        return t

    def page_break():
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Page header and footer on every page: the report title; Confidential and page N of M.
    sec.header.paragraphs[0].text = m["title"]
    foot = sec.footer.paragraphs[0]
    foot.add_run("Confidential, page ")
    field(foot, "PAGE", "1")
    foot.add_run(" of ")
    field(foot, "NUMPAGES", "1")

    # Cover
    doc.add_heading(m["title"], level=0)
    table(m["cover"])
    doc.add_paragraph(m["confidential"])
    page_break()

    # Contents: a Word TOC field over the headings below.
    p = doc.add_paragraph()
    p.add_run("Contents").bold = True
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
          "Open in Word and update fields to build the contents.")
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)
    page_break()

    for block in m["blocks"]:
        kind = block[0]
        if kind == "h1":
            doc.add_heading(block[1], level=1)
        elif kind in ("h2", "group"):
            doc.add_heading(block[1], level=2)
        elif kind == "table":
            table(block[1], header=block[2])
        elif kind == "para":
            doc.add_paragraph(block[1])
        elif kind == "counts":
            table(block[1], header=block[2])
            for term, meaning in block[3]:
                p = doc.add_paragraph()
                p.add_run(term + ": ").bold = True
                p.add_run(meaning)
            doc.add_paragraph().add_run(block[4]).bold = True
        elif kind == "ready":
            doc.add_paragraph().add_run(block[1]).bold = True
        elif kind == "case":
            r, values, response, rows, steps, unread = block[1:]
            doc.add_heading(r["id"] + ": " + r["title"], level=3)
            table([[k, v] for k, v in zip(CASE_FIELDS, values)])
            if steps:
                table(steps, header=["Step", "Result", "Note"])
            if response:
                p = doc.add_paragraph()
                p.add_run("Response with code: ").bold = True
                code = p.add_run(response)
                code.font.name = "Courier New"
                code.font.size = Pt(9)
            for note in unread:
                doc.add_paragraph(note)
            text_mm = width / 36000.0
            for row in rows:
                # The PDF's layout: a landscape screenshot across the page, portrait ones three to a row.
                if row[0]["kind"] == "tall":
                    grid = doc.add_table(rows=1, cols=len(row))
                    for cell, shot in zip(grid.rows[0].cells, row):
                        crop, w, h = shot_parts(shot, text_mm)[0]
                        cell.paragraphs[0].add_run().add_picture(embedded(shot, crop, w, h), width=Mm(w), height=Mm(h))
                        cell.add_paragraph(caption(shot, 1, 1, False), style="Caption")
                    continue
                for shot in row:
                    parts = shot_parts(shot, text_mm)
                    for n, (crop, w, h) in enumerate(parts, 1):
                        doc.add_picture(embedded(shot, crop, w, h), width=Mm(w), height=Mm(h))
                        doc.add_paragraph(caption(shot, n, len(parts), False), style="Caption")

    # Mask every run of text the document holds, tables, header and footer included.
    def paragraphs(container):
        for p in container.paragraphs:
            yield p
        for t in container.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in paragraphs(cell):
                        yield p
    for part in (doc, sec.header, sec.footer):
        for p in paragraphs(part):
            for r in p.runs:
                if r.text and r.text not in m["safe"]:
                    masked = mask(r.text)
                    if masked != r.text:
                        r.text = masked
    doc.save(str(out))
    print("wrote " + str(out))


# The ABDM Sandbox design system: colours, and the fonts it names with their PDF fallbacks.
INK = {"navy": "#112044", "deep": "#0B1630", "blue": "#264488", "ink": "#1C2333", "ink2": "#3A4460",
       "ink3": "#4A5368", "ink4": "#6A738A", "orange": "#E8791F", "otext": "#B4530F", "peach": "#FBF1E7",
       "mist": "#F4F6FA", "line": "#DDE2EE", "dark": "#1C2333", "ice": "#C9D4EE", "off": "#FFFDF8"}
FONTS = {"serif": (["newsreader-light.ttf", "newsreader_18pt-light.ttf", "newsreader-regular.ttf", "newsreader.ttf"],
                   "Times-Roman", "Newsreader"),
         "sans": (["publicsans-regular.ttf", "publicsans.ttf"], "Helvetica", "Public Sans"),
         "mono": (["jetbrainsmono-regular.ttf", "jetbrainsmono.ttf"], "Courier", "JetBrains Mono")}
MILESTONES = {"M1": "M1 Identity", "M2": "M2 Health Information Provider", "M3": "M3 Health Information User",
              "M4": "M4 Registry Integration"}
# Column widths, as shares of the text width, for the tables the report holds.
WIDTHS = {
    ("Version", "Date", "By", "Change"): [.12, .18, .25, .45],
    ("Name", "Organisation", "Role"): [.34, .36, .30],
    ("Measure", "All cases", "Mandatory"): [.56, .22, .22],
    ("Function", "Cases", "Passed", "Conditional pass", "Failed", "Blocked", "Not tested", "Not applicable"):
        [.36, .07, .08, .11, .08, .09, .10, .11],
    ("Test case ID", "Title", "Condition"): [.24, .30, .46],
    ("Step", "Result", "Note"): [.10, .22, .68],
    ("Test case ID", "Title", "Marking", "Status", "Needed, or why not applicable"): [.22, .30, .14, .12, .22],
    ("Test case ID", "Title", "Agreed by and when"): [.25, .37, .38],
    ("Test case ID", "Status", "Iterations", "Failure reason", "Error code"): [.22, .12, .11, .37, .18],
    ("Flow", "ABHA address used", "HI types created or tested"): [.30, .35, .35],
    ("Name", "Role", "Qualifications"): [.34, .26, .40],
    ("Sign-off", "Name", "Date", "Signature"): [.18, .34, .18, .30],
}


def find_fonts(extra=None):
    """Register Newsreader, Public Sans and JetBrains Mono where found, else Times-Roman, Helvetica and Courier."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    dirs = ([Path(extra)] if extra else []) + [
        Path("abdm-ft/fonts"), Path(__file__).resolve().parent / "fonts", Path.home() / "Library" / "Fonts",
        Path("/Library/Fonts"), Path("/usr/share/fonts")]
    chosen, notes = {}, []
    for role, (names, fallback, family) in FONTS.items():
        for d in dirs:
            if role in chosen or not d.is_dir():
                continue
            files = {}
            for f in d.rglob("*.ttf"):
                files.setdefault(f.name.lower(), f)
            for name in names:
                if name in files and role not in chosen:
                    try:
                        pdfmetrics.registerFont(TTFont("abdm-" + role, str(files[name])))
                        chosen[role] = "abdm-" + role
                        notes.append(family + " from " + str(files[name]))
                    except Exception:
                        pass
        if role not in chosen:
            chosen[role] = fallback
            notes.append(family + " not found, " + fallback + " in its place")
    print("fonts: " + "; ".join(notes))
    return chosen


def write_pdf(m, out, fonts_dir=None):
    """The report to submit, drawn in the ABDM Sandbox design system with ReportLab."""
    from io import BytesIO
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import (BaseDocTemplate, CondPageBreak, Flowable, Frame, Image as Picture,
                                    NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table,
                                    TableStyle)

    F = find_fonts(fonts_dir)
    C = {k: colors.HexColor(v) for k, v in INK.items()}
    W, H = A4
    M = 20 * mm
    TW = W - 2 * M
    safe = m["safe"]

    def clean(text):
        """Every string drawn goes through the DOCX's mask, a list of UUID REQUEST-IDs excepted."""
        text = str(text if text is not None else "")
        return text if text in safe else mask(text)

    def style(name, font, size, lead, color, **kw):
        return ParagraphStyle(name, fontName=F[font], fontSize=size, leading=lead, textColor=C[color], **kw)

    S = {
        "body": style("body", "sans", 9.5, 14.5, "ink3", spaceAfter=6),
        "lead": style("lead", "sans", 11, 17, "ink2", spaceAfter=8),
        "cell": style("cell", "sans", 8.8, 12.6, "ink3"),
        "strong": style("strong", "sans", 8.8, 12.6, "ink"),
        "id": style("id", "mono", 7.8, 11.5, "ink"),
        "code": style("code", "mono", 7.6, 11, "ink2"),
        "h1": style("h1", "serif", 22, 26, "navy", spaceAfter=10),
        "h2": style("h2", "serif", 15, 19, "navy", spaceBefore=12, spaceAfter=6),
        "case": style("case", "serif", 12, 15, "navy", spaceAfter=4),
        "small": style("small", "sans", 8.2, 11, "ink"),
        "term": style("term", "sans", 8.2, 11.6, "ink3"),
        "note": style("note", "sans", 9.5, 14.5, "ink2"),
        "chev": style("chev", "serif", 13, 14.5, "orange"),
        "page": style("page", "mono", 8, 12, "ink4", alignment=TA_RIGHT),
        "toc": style("toc", "sans", 10.5, 14, "ink"),
        "tocpage": style("tocpage", "mono", 7.2, 10, "ink4", alignment=TA_RIGHT),
        "digit": style("digit", "mono", 24, 28, "off"),
        "cover": style("cover", "serif", 38, 43, "off"),
        "coversub": style("coversub", "sans", 11.5, 16, "ice"),
        "fact": style("fact", "sans", 10, 13.5, "ink"),
    }

    def para(text, s="body"):
        return Paragraph(escape(clean(text)), S[s])

    class Tracked(Flowable):
        """Mono uppercase with wide tracking, which a Paragraph cannot draw: wraps on spaces."""

        def __init__(self, text, size=7.2, color="otext", space=None, upper=True, after=0):
            Flowable.__init__(self)
            text = clean(text)
            self.text = text.upper() if upper else text
            self.size, self.color, self.after = size, color, after
            self.space = size * 0.12 if space is None else space
            self.lead = size * 1.45

        def measure(self, s):
            return stringWidth(s, F["mono"], self.size) + self.space * max(len(s) - 1, 0)

        def wrap(self, aw, ah):
            lines, cur = [], ""
            for word in self.text.split(" "):
                trial = cur + " " + word if cur else word
                if cur and self.measure(trial) > aw:
                    lines.append(cur)
                    cur = word
                else:
                    cur = trial
            lines.append(cur)
            widest = max(self.measure(x) for x in lines)
            if widest > aw:
                # One word wider than its cell, such as ITERATIONS: tighten the tracking, then the size.
                self.space = 0
                widest = max(self.measure(x) for x in lines)
                if widest > aw:
                    self.size *= aw / widest
                    widest = aw
            self.lines = lines
            self.width = min(aw, widest)
            self.height = len(lines) * self.lead + self.after
            return self.width, self.height

        def draw(self):
            c = self.canv
            c.setFont(F["mono"], self.size)
            c.setFillColor(C[self.color])
            for i, line in enumerate(self.lines):
                # The first baseline sits where a 9pt Paragraph's would, so labels line up with cell text.
                c.drawString(0, self.height - self.after - i * self.lead - self.size * 1.2, line, charSpace=self.space)

    class Rule(Flowable):
        """A rule across the text width. At the top of a page it drops the space above it, and a
        hairline between cases drops altogether, so it never doubles the header's rule."""

        def __init__(self, color="dark", width=1, before=0, after=10, top=True):
            Flowable.__init__(self)
            self.color, self.weight, self.before, self.after, self.top = color, width, before, after, top

        def wrap(self, aw, ah):
            self.width = aw
            self.skip = ah >= H - 2 * M - 1 and not self.top
            at_top = ah >= H - 2 * M - 1
            self.height = 0 if self.skip else (0 if at_top else self.before) + self.after + self.weight
            return aw, self.height

        def draw(self):
            if self.skip:
                return
            self.canv.setStrokeColor(C[self.color])
            self.canv.setLineWidth(self.weight)
            y = self.after + self.weight / 2.0
            self.canv.line(0, y, self.width, y)

    class Mark(Flowable):
        """A zero size marker whose page the contents list records."""

        def __init__(self, key):
            Flowable.__init__(self)
            self.toc_key = key

        def wrap(self, aw, ah):
            return 0, 0

        def draw(self):
            pass

    def status(text):
        word = clean(text).split(" ")[0].upper()
        color = "blue" if word in ("PASS", "CONDITIONAL") else "otext" if word == "FAIL" else "ink4"
        return Tracked(text or "Not recorded", size=7.2, color=color, space=0.6)

    def grid(rows, header=None, pad=4.2, value="strong"):
        """Hairline rules only, a mist header row of mono labels that repeats, Paragraphs that wrap."""
        cols = len(header or rows[0])
        shares = WIDTHS.get(tuple(header or ())) or ([.34, .66] if cols == 2 else [1.0 / cols] * cols)
        body = []
        split = []
        for row in rows:
            # A cell too long for one page goes on over rows, so a table never outgrows a page.
            texts = shown(row)
            chunks = [pieces_of(x) for x in texts]
            for n in range(max(len(c) for c in chunks)):
                split.append(([c[n] if n < len(c) else "" for c in chunks], n))
        for texts, n in split:
            cells = []
            for c, text in enumerate(texts):
                name = header[c] if header else ("Label" if c == 0 else texts[0])
                if n and (c == 0 or not text):
                    cells.append("")
                elif not header and c == 0:
                    cells.append(Tracked(text, size=6.8, color="ink4", space=0.7))
                elif name == "Status":
                    cells.append(status(text))
                elif name == "Test case ID":
                    cells.append(para(text, "id"))
                elif not header:
                    cells.append(para(text, value))
                else:
                    cells.append(para(text, "cell"))
            body.append(cells)
        data = ([[Tracked(h, size=6.6, color="ink3", space=0.7) for h in header]] if header else []) + body
        t = Table(data, colWidths=[TW * s for s in shares], repeatRows=1 if header else 0, hAlign="LEFT",
                  rowHeights=None)
        cmds = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 7),
                ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), pad),
                ("BOTTOMPADDING", (0, 0), (-1, -1), pad),
                ("LINEBELOW", (0, 0), (-1, -1), 0.6, C["line"]), ("LINEABOVE", (0, 0), (-1, 0), 0.6, C["line"])]
        if header:
            cmds += [("BACKGROUND", (0, 0), (-1, 0), C["mist"]), ("TOPPADDING", (0, 0), (-1, 0), 7),
                     ("BOTTOMPADDING", (0, 0), (-1, 0), 5.5)]
            if header[0] == "Sign-off":
                cmds += [("TOPPADDING", (0, 1), (-1, -1), 14), ("BOTTOMPADDING", (0, 1), (-1, -1), 14)]
        t.setStyle(TableStyle(cmds))
        return t

    def callout(lines, s="note"):
        rows = [[Paragraph("›", S["chev"]), para(line, s)] for line in lines]
        t = Table(rows, colWidths=[28, TW - 28], hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), C["peach"]), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (0, -1), 13), ("LEFTPADDING", (1, 0), (1, -1), 2),
            ("RIGHTPADDING", (0, 0), (0, -1), 0), ("RIGHTPADDING", (1, 0), (1, -1), 16), ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5), ("TOPPADDING", (0, 0), (-1, 0), 12),
            ("BOTTOMPADDING", (0, -1), (-1, -1), 12)]))
        return t


    def facts(values):
        """A tested case's fields, each a labelled line or paragraph, hairlines between them. The
        result and the REQUEST-IDs share the first line; long text goes on over rows."""
        def cell(name, text, first=True):
            label = ('<font name="' + F["mono"] + '" size="6.2" color="' + INK["ink4"] + '">' +
                     escape(name.upper()) + "</font>&nbsp;&nbsp; ") if first else ""
            return Paragraph(label + escape(clean(text)), S["code" if name == "Request IDs" else "small"])

        rows = [[cell(CASE_FIELDS[0], values[0]), cell(CASE_FIELDS[3], values[3])]]
        spans = []
        for name, text in zip(CASE_FIELDS[1:3], values[1:3]):
            for n, piece in enumerate(pieces_of(text, 1200)):
                spans.append(("SPAN", (0, len(rows)), (1, len(rows))))
                rows.append([cell(name, piece, n == 0), ""])
        t = Table(rows, colWidths=[TW * 0.56, TW * 0.44], hAlign="LEFT")
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                               ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 1.6),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6), ("LINEBELOW", (0, 0), (-1, -1), 0.6, C["line"]),
                               ("RIGHTPADDING", (0, 0), (0, 0), 10)] + spans))
        return t

    def terms(rows):
        """Each term of the counts table on one line: the term, then what it means."""
        return [Paragraph("<b>" + escape(clean(term)) + "</b>: " + escape(clean(meaning)), S["term"])
                for term, meaning in rows] + [Spacer(1, 8)]

    def framed(shot, crop, w, h, text):
        pic = Picture(embedded(shot, crop, w, h), width=w * mm, height=h * mm)
        frame = Table([[pic]], colWidths=[w * mm], rowHeights=[h * mm], hAlign="LEFT")
        frame.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, C["line"]),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                   ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        return [frame, Spacer(1, 2), Tracked(text, size=6.2, color="ink4", space=0, upper=False)]

    def shots(row):
        """One row of screenshots as flowables: portrait ones side by side; a landscape one alone;
        a very tall one as one flowable per piece, so each piece fits a page."""
        out = []
        if row[0]["kind"] == "tall":
            cells = []
            for shot in row:
                crop, w, h = shot_parts(shot, TW / mm)[0]
                cells.append(framed(shot, crop, w, h, caption(shot, 1, 1)))
            out.append(Table([cells], colWidths=[TW / 3.0] * len(row), hAlign="LEFT"))
        else:
            for shot in row:
                parts = shot_parts(shot, TW / mm)
                for n, (crop, w, h) in enumerate(parts, 1):
                    out.append(Table([[framed(shot, crop, w, h, caption(shot, n, len(parts)))]], colWidths=[TW],
                                     hAlign="LEFT"))
        for f in out:
            f.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        return out

    def opener(number, name):
        return [CondPageBreak(120), Rule("dark", 1, before=12, after=10), Mark("s" + str(number)),
                Tracked("%02d · %s" % (number, name), size=7.5, color="otext", space=1.0, after=8),
                Paragraph(escape(clean(name)), S["h1"])]

    def story(marks, total):
        out = [NextPageTemplate("inner"), PageBreak()]
        out += [Rule("dark", 1, after=12), Tracked("Contents", size=7.5, color="otext", space=1.0, after=8),
                Paragraph("Contents", S["h1"])]
        functions = [b[1] for b in m["blocks"] if b[0] == "group"]

        def sections(names, first):
            rows = [[Tracked("%02d" % (first + i), size=7.5, color="otext", space=0.6), para(name, "toc"),
                     para(str(marks.get("s" + str(first + i), 0)), "page")] for i, name in enumerate(names)]
            t = Table(rows, colWidths=[40, TW - 80, 40], hAlign="LEFT")
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 4),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                                   ("LINEBELOW", (0, 0), (-1, -1), 0.6, C["line"])]))
            return t

        split = SECTIONS.index("Evidence") + 1
        out.append(sections(SECTIONS[:split], 1))
        if functions:
            # Under Evidence, each function group with the page it starts on.
            rows = [["", para(name, "small"), para(str(marks.get("g" + name, 0)), "tocpage")] for name in functions]
            groups_t = Table(rows, colWidths=[40, TW - 80, 40], hAlign="LEFT")
            groups_t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                          ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                                          ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5), ("TOPPADDING", (0, 0), (-1, 0), 6),
                                          ("BOTTOMPADDING", (0, -1), (-1, -1), 6),
                                          ("LINEBELOW", (0, -1), (-1, -1), 0.6, C["line"])]))
            out.append(groups_t)
        out.append(sections(SECTIONS[split:], split + 1))
        out.append(PageBreak())

        # A heading never ends a page alone: a section opener, a group heading or a case heading is
        # held, then kept on one page with the first thing after it.
        section, held = 0, []

        def place(lead, rest=()):
            # Ask for the room the held headings and the lead take, with a margin; a lead taller than
            # a page asks for the headings and its first lines.
            body = H - 2 * M
            need = sum(f.wrap(TW, body)[1] + f.getSpaceBefore() + f.getSpaceAfter() for f in held + list(lead))
            if need > body:
                need = sum(f.wrap(TW, body)[1] for f in held) + 80
            out.append(CondPageBreak(need + 14))
            out.extend(held + list(lead) + list(rest))
            del held[:]

        for block in m["blocks"]:
            kind = block[0]
            if kind == "h1":
                section += 1
                parts = opener(section, block[1])
                out.append(parts[0])
                held.extend(parts[1:])
            elif kind == "h2":
                held.append(Paragraph(escape(clean(block[1])), S["h2"]))
            elif kind == "group":
                held.extend([Mark("g" + block[1]), Paragraph(escape(clean(block[1])), S["h2"])])
            elif kind == "table":
                place([grid(block[1], block[2]), Spacer(1, 10)])
            elif kind == "counts":
                # The counts, what each means and the verdict stay on one page.
                place([grid(block[1], block[2], pad=3), Spacer(1, 6)] + terms(block[3]) +
                      [callout([block[4]], "lead"), Spacer(1, 10)])
            elif kind == "ready":
                place([callout([block[1]], "lead"), Spacer(1, 10)])
            elif kind == "para":
                place([para(block[1])])
            elif kind == "case":
                # A case flows on from the one before; its heading and fields stay with its first screenshot.
                r, values, response, rows, steps, unread = block[1:]
                lead = [Rule("line", 0.6, before=2, after=6, top=False), Mark("c" + r["id"]),
                        Paragraph('<font name="' + F["mono"] + '" size="8" color="' + INK["otext"] + '">' +
                                  escape(clean(r["id"])) + "</font>&nbsp;&nbsp; " + escape(clean(r["title"])),
                                  S["case"]), facts(values)]
                lead += [para(note, "small") for note in unread]
                if steps:
                    lead += [Spacer(1, 4), grid(steps, ["Step", "Result", "Note"], pad=2.6)]
                if response:
                    code = Table([[Tracked("Response with code", size=6.6, color="ink4", space=0.7, after=3)]] +
                                 [[para(piece, "code")] for piece in pieces_of(response, 1500)],
                                 colWidths=[TW], hAlign="LEFT")
                    code.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), C["mist"]),
                                              ("LEFTPADDING", (0, 0), (-1, -1), 9), ("RIGHTPADDING", (0, 0), (-1, -1), 9),
                                              ("TOPPADDING", (0, 0), (-1, 0), 6), ("BOTTOMPADDING", (0, -1), (-1, -1), 6),
                                              ("TOPPADDING", (0, 1), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -2), 0)]))
                    lead += [Spacer(1, 4), code]
                pictures = [f for row in rows for f in shots(row)]
                place(lead + pictures[:1], pictures[1:] + [Spacer(1, 6)])
        if held:
            place([])
        return out

    eyebrow = clean(m["title"]).upper()

    def tracked(c, x, y, text, size, color, space, right=False):
        text = clean(text)
        c.setFont(F["mono"], size)
        c.setFillColor(C[color])
        if right:
            x -= stringWidth(text, F["mono"], size) + space * (len(text) - 1)
        c.drawString(x, y, text, charSpace=space)

    def inner(c, d):
        c.saveState()
        tracked(c, M, H - 13 * mm, eyebrow, 6.8, "ink4", 0.8)
        c.setStrokeColor(C["line"])
        c.setLineWidth(0.6)
        c.line(M, H - 15.5 * mm, W - M, H - 15.5 * mm)
        c.line(M, 15 * mm, W - M, 15 * mm)
        tracked(c, M, 11 * mm, "CONFIDENTIAL", 6.8, "ink4", 0.8)
        tracked(c, W - M, 11 * mm, "PAGE " + str(d.page) + " OF " + str(d.total or d.page), 6.8, "ink4", 0.8,
                right=True)
        c.restoreState()

    def cover(c, d):
        import math
        c.saveState()
        top = H * 0.47
        c.setFillColor(C["navy"])
        c.rect(0, top, W, H - top, stroke=0, fill=1)
        # The Ashoka wheel, bleeding off the top right: 24 spokes and three rings, ice at 11 percent.
        p = c.beginPath()
        p.rect(0, top, W, H - top)
        c.clipPath(p, stroke=0, fill=0)
        cx, cy, R = W - 22 * mm, H - 18 * mm, 92 * mm
        c.setStrokeColor(C["ice"])
        c.setStrokeAlpha(0.11)
        c.setLineWidth(1.4)
        for radius in (R, R * 0.86, R * 0.16):
            c.circle(cx, cy, radius, stroke=1, fill=0)
        for k in range(24):
            a = 2 * math.pi * k / 24
            c.line(cx + R * 0.16 * math.cos(a), cy + R * 0.16 * math.sin(a),
                   cx + R * 0.86 * math.cos(a), cy + R * 0.86 * math.sin(a))
        c.restoreState()
        c.saveState()
        c.setFillColor(C["orange"])
        c.rect(M, H - 26 * mm, 18 * mm, 2, stroke=0, fill=1)
        tracked(c, M, H - 33 * mm, "ABDM SANDBOX · FUNCTIONAL TEST REPORT", 7.8, "orange", 1.4)
        if m.get("draft"):
            tracked(c, M, H - 39 * mm, "DRAFT · " + str(len(m["draft"])) + " MANDATORY CASES NOT YET PASS", 7.8, "ice", 1.4)
        tail = MILESTONES.get(m["milestone"].upper(), m["milestone"])
        lead = ("Draft " + m["label"]) if m.get("draft") else (m["label"][:1].upper() + m["label"][1:])
        head = (escape(clean(lead + " report, ")) +
                '<font color="' + INK["orange"] + '">' + escape(clean(tail)) + "</font>")
        title = Paragraph(head, S["cover"])
        _, th = title.wrap(TW * 0.86, H)
        sub = Paragraph(escape(clean(m["product"])) + " &#183; " + escape(clean(m["organisation"])), S["coversub"])
        _, sh = sub.wrap(TW * 0.86, H)
        sub.drawOn(c, M, top + 16 * mm)
        title.drawOn(c, M, top + 16 * mm + sh + 8 * mm)
        # Key facts as a ruled-box grid: four across, then three.
        cells = []
        for label, value in m["facts"]:
            cells.append([Tracked(label, size=6.6, color="ink4", space=0.7, after=4), para(value, "fact")])
        rows = [cells[:4], [cells[4], cells[6], cells[5], ""]]
        facts = Table(rows, colWidths=[TW / 4.0] * 4)
        facts.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, C["line"]), ("INNERGRID", (0, 0), (-1, -1), 0.8, C["line"]),
                                   ("SPAN", (2, 1), (3, 1)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                                   ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 12)]))
        _, fh = facts.wrap(TW, H)
        facts.drawOn(c, M, top - 14 * mm - fh)
        note = callout([m["confidential"]])
        _, nh = note.wrap(TW, H)
        note.drawOn(c, M, 22 * mm)
        c.restoreState()

    class Doc(BaseDocTemplate):
        def afterFlowable(self, f):
            key = getattr(f, "toc_key", None)
            if key and key not in self.marks:
                self.marks[key] = self.page

    marks, total, pdf = {}, 0, b""
    for attempt in range(4):
        buf = BytesIO()
        d = Doc(buf, pagesize=A4, leftMargin=M, rightMargin=M, topMargin=M, bottomMargin=M,
                title=clean(m["title"]), author=clean(m["organisation"]), subject="ABDM functional test report")
        d.marks, d.total = {}, total
        frame = Frame(M, M, TW, H - 2 * M, id="body", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
        d.addPageTemplates([PageTemplate("cover", [frame], onPage=cover), PageTemplate("inner", [frame], onPage=inner)])
        d.build(story(marks, total))
        pdf = buf.getvalue()
        if d.marks == marks and d.page == total:
            break
        marks, total = d.marks, d.page
    Path(out).write_bytes(pdf)
    print("wrote " + str(out))


def build(path, run, meta, evidence, out, recs, groups, manifest=None, fmt="both", fonts_dir=None, draft=None):
    """Write the report: the PDF to submit and the DOCX to edit, beside each other."""
    m = content(path, run, meta, evidence, recs, groups, manifest, draft, out)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    if m["figures"]:
        # The full size screenshots go in the annexure bundle; this lists them figure by figure.
        listing = out.with_name(out.stem + "-screenshots.json")
        listing.write_text(json.dumps({"report": out.stem, "figures": m["figures"]}, indent=1))
        print("wrote " + str(listing))
    if fmt in ("docx", "both"):
        write_docx(m, out.with_suffix(".docx"))
    if fmt in ("pdf", "both"):
        write_pdf(m, out.with_suffix(".pdf"), fonts_dir)


# A module skill built per functionality ships functionalities.json beside this
# script: each functionality's cases and their titles, and the screen requirements
# in its file with the cases each cites and the functionality it hands a case to.
# The tracker and the findings are written from it and from the workbook; the
# markings stay in the workbook, read as plan reads them, and a case is through
# by the same passes_gate the gate uses.
TRACKER_SET = {"building": "building", "built": "built and waiting for the tester", "tested": "tested"}
HELD = ("building", "built and waiting for the tester")


def load_functionalities(path=None):
    p = Path(path) if path else Path(__file__).resolve().parent / "functionalities.json"
    if not p.exists():
        return None
    data = json.loads(p.read_text())
    return data if data.get("functionalities") else None


def shown_result(r, agreed=()):
    """A case's result as the tracker and the findings show it."""
    if r is None:
        return "not in the workbook"
    if r["status"] == "CONDITIONAL PASS" and not passes_gate(r, agreed):
        return "CONDITIONAL PASS with no condition"
    if r["status"] == "NA" and r["id"] in agreed:
        return "NA, agreed"
    if r["status"] in ("PASS", "CONDITIONAL PASS", "FAIL", "NA"):
        return r["status"]
    if not r["status"] and r["remarks"].lower().startswith("needs tester"):
        return "needs tester"
    return r["status"] or "open"


def run_date(path):
    """The Test Execution date the writer put in the sheet's header block, when it holds one."""
    wb = load(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=15):
            for cell in row:
                text = norm(cell.value)
                if text.startswith("Test Execution date:"):
                    m = re.search(r"\d{4}-\d{2}-\d{2}", text)
                    if m:
                        return m.group(0)
    return ""


def judged_on(f, by_id, state):
    """A functionality's scope and the cases it is judged on: each mandatory case alone, an either pair
    together, or every applicable case of an optional functionality the tester added."""
    present = [by_id[c] for c in f["cases"] if c in by_id]
    if present and all(r.get("na") for r in present):
        return "not applicable", []
    mandatory, seen = [], set()
    for r in present:
        if r["kind"] == "mandatory":
            mandatory.append([r["id"]])
        elif r["kind"] == "pair":
            ids = tuple(i for i in r["pair"] if i in by_id)
            if ids not in seen:
                seen.add(ids)
                mandatory.append(list(ids))
    if mandatory:
        return "mandatory", mandatory
    if f["id"] in state.get("in_scope", []):
        return "optional, added", [[r["id"]] for r in present if not r.get("na")]
    return "optional", []


def from_workbook(groups_of, by_id, agreed, tested):
    """The state the workbook gives: passed, failing, built and waiting for the tester, or not started."""
    def through(ids):
        return any(passes_gate(by_id[i], agreed) for i in ids)

    if all(through(ids) for ids in groups_of):
        return "passed"
    if any(by_id[i]["status"] == "FAIL" for ids in groups_of if not through(ids) for i in ids):
        return "failing"
    if tested or any(shown_result(by_id[i], agreed) != "open" for ids in groups_of for i in ids):
        return "built and waiting for the tester"
    return "not started"


def run_tracker(a, entity, agreed):
    """Write the functionality tracker into INTEGRATION-PLAN.md from the workbook, functionalities.json
    and the tracker state, and print where each functionality stands, by id."""
    data = load_functionalities(a.functionalities)
    if not data:
        sys.exit("no functionalities.json beside this script: this module's skill is not built per functionality "
                 "yet, so it has no tracker. Use its INTEGRATION-PLAN.md steps instead")
    state_path = Path(a.state) if a.state else Path(a.workbook).resolve().parent / "tracker.json"
    state = json.loads(state_path.read_text()) if state_path.exists() else {}
    entity = entity or state.get("entity")
    if entity not in ("government", "private"):
        sys.exit("the tracker holds no entity type: ask the tester, government or private, then pass --entity")
    if state.get("entity") and state["entity"] != entity:
        sys.exit("the tracker says " + state["entity"] + " and --entity says " + entity + ": ask the tester which")
    state["entity"], state["module"] = entity, data.get("module", "")
    by_fid = {f["id"]: f for f in data["functionalities"]}
    owner = {c: f["id"] for f in data["functionalities"] for c in f["cases"]}
    recs, _ = read_cases(a.workbook, a.run, entity)
    by_id = {r["id"]: r for r in recs}
    today = date.today().isoformat()

    def snapshot(fid):
        """What the workbook holds for a functionality's cases now, to tell when a later run changed it."""
        return {c: shown_result(by_id.get(c), agreed) for c in by_fid[fid]["cases"]}

    for item in a.set or []:
        fid, _, what = item.partition("=")
        if fid not in by_fid or what not in list(TRACKER_SET) + ["in-scope"]:
            sys.exit("--set takes <functionality id>=building|built|tested|in-scope, not " + item)
        if what == "in-scope":
            state["in_scope"] = sorted(set(state.get("in_scope", [])) | {fid})
            continue
        entry = state.setdefault("functionalities", {}).setdefault(fid, {})
        entry.update(state=TRACKER_SET[what], date=today, snapshot=snapshot(fid))
        if a.build:
            entry["build"] = mask(norm(a.build))
    for item in a.screen or []:
        case, _, where = item.partition("=")
        if case not in owner or not where.strip():
            sys.exit("--screen takes <case id>=<screen or route> for a case in functionalities.json, not " + case)
        state.setdefault("screens", {})[case] = mask(norm(where))
    when = run_date(a.workbook)
    rows, counts, tally_of = [], {}, {"passed": 0, "failing": 0, "open": 0, "total": 0}
    for f in data["functionalities"]:
        scope, groups_of = judged_on(f, by_id, state)
        entry = state.get("functionalities", {}).get(f["id"], {})
        if scope == "not applicable":
            st = "not applicable"
        elif scope == "optional":
            st = "optional, not in scope"
        else:
            st = from_workbook(groups_of, by_id, agreed, entry.get("state") == "tested")
            if entry.get("state") in HELD:
                # A held state gives way once a later write into the workbook settles the functionality.
                if st in ("passed", "failing") and entry.get("snapshot") != snapshot(f["id"]):
                    entry.pop("state")
                    entry.pop("snapshot", None)
                else:
                    st = entry["state"]
        listed = []
        for ids in groups_of:
            listed.append(" or ".join(i + " " + shown_result(by_id.get(i), agreed) for i in ids))
            tally_of["total"] += 1
            if any(passes_gate(by_id[i], agreed) for i in ids if i in by_id):
                tally_of["passed"] += 1
            elif any(by_id[i]["status"] == "FAIL" for i in ids if i in by_id):
                tally_of["failing"] += 1
            else:
                tally_of["open"] += 1
        screens = "; ".join(i + ": " + state.get("screens", {}).get(i, "not mapped") for ids in groups_of for i in ids)
        tested = any(shown_result(by_id.get(i), agreed) not in ("open", "not in the workbook") for ids in groups_of for i in ids)
        rows.append([str(f["order"]), f["title"] + " (`" + f["id"] + "`)",
                     ", ".join(listed) or ("none for a " + entity + " entity"), screens, st,
                     entry.get("date") or (when if tested else ""), entry.get("build", "not recorded" if tested else "")])
        if scope not in ("not applicable", "optional"):
            counts[st] = counts.get(st, 0) + 1
        failing = [i for ids in groups_of for i in ids if shown_result(by_id.get(i), agreed) == "FAIL"]
        print(f["id"] + ": " + st + (" (" + ", ".join(failing) + " FAIL)" if failing else ""))
    in_scope = sum(counts.values())
    module = data.get("module", "").upper()
    lines = ["<!-- abdm-tracker start: written by abdm_report.py tracker. Never edit it by hand -->",
             "",
             "Entity type: " + entity + ". Run: " + a.run + ". Workbook: " + Path(a.workbook).name + ". Written " + today + ".",
             "",
             "In scope: " + str(in_scope) + " functionalities. " +
             ", ".join(str(counts[k]) + " " + k for k in ("passed", "failing", "built and waiting for the tester",
                                                         "building", "not started") if counts.get(k)) + ". " +
             "Mandatory cases: " + str(tally_of["total"]) + ", " + str(tally_of["passed"]) + " passed, " +
             str(tally_of["failing"]) + " failing, " + str(tally_of["open"]) + " open.",
             "",
             "| # | Functionality | Mandatory cases and results | Screen or route | State | Date | Build |",
             "|---|---|---|---|---|---|---|"]
    lines += ["| " + " | ".join(c.replace("|", "/") for c in row) + " |" for row in rows]
    lines += ["", "<!-- abdm-tracker end -->"]
    block = "\n".join(lines)
    plan = Path(a.plan)
    text = plan.read_text() if plan.exists() else "# Integration plan\n"
    found = re.search(r"<!-- abdm-tracker start.*?<!-- abdm-tracker end -->", text, re.S)
    if found:
        text = text[:found.start()] + block + text[found.end():]
    else:
        text = text.rstrip("\n") + "\n\n## Functionality tracker, " + module + "\n\n" + block + "\n"
    plan.write_text(text)
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_path.write_text(json.dumps(state, indent=1))
    print("in scope " + str(in_scope) + "; mandatory cases " + str(tally_of["total"]) + ", " +
          str(tally_of["passed"]) + " passed, " + str(tally_of["failing"]) + " failing, " +
          str(tally_of["open"]) + " open")
    print("wrote " + str(plan) + " and " + str(state_path))


def words(text):
    return set(re.findall(r"[a-z0-9]+", str(text).lower())) - {"the", "a", "an", "of", "and", "or", "to", "is"}


def coverage(case, f, by_fid):
    """How a functionality's file covers one case: ("own", rule) when a screen requirement states its check
    for the case; ("through", rule, target, counterpart) when the only rows citing it hand it to another
    functionality and that one states a check for the case of the same title; else ("none", reason)."""
    if not f:
        return ("none", "the case is in no functionality")
    rows = [q for q in f.get("screen_requirements", []) if case in q["cases"]]
    own = [q for q in rows if not q.get("delegates_to")]
    if own:
        return ("own", own[0]["requirement"], len(own))
    title = words(f.get("titles", {}).get(case, ""))
    for q in rows:
        for target in q.get("delegates_to", []):
            t = by_fid.get(target)
            if not t:
                continue
            # The case of the same title, else the closest; a tie goes to the first in the file.
            scored = sorted(((-(len(title & words(tt)) / max(1, len(title | words(tt)))), i, c)
                             for i, (c, tt) in enumerate(t.get("titles", {}).items())))
            scored = [(-x, c) for x, _, c in scored]
            if not scored or scored[0][0] < 0.5:
                continue
            counterpart = scored[0][1]
            rules = [r for r in t.get("screen_requirements", []) if counterpart in r["cases"] and not r.get("delegates_to")]
            if rules:
                return ("through", rules[0]["requirement"], target, counterpart)
    if rows:
        return ("none", "it hands the case to " + ", ".join(sorted({t for q in rows for t in q["delegates_to"]})) +
                ", which states no check for its case of the same title")
    return ("none", "no screen requirement cites it")


def cut(text, size=120):
    return text if len(text) <= size else text[:size - 3].rsplit(" ", 1)[0] + "..."


def write_findings(path, run, entity, out, data, recs=None):
    """<report>-findings.md beside the report: each FAIL and CONDITIONAL PASS, its functionality, and the screen
    requirement that states the check for it, in its own file or in the one it hands the case to. A case with
    none is a Catalogue issue to file."""
    if recs is None:
        recs, _ = read_cases(path, run, entity)
    by_fid = {f["id"]: f for f in data["functionalities"]}
    owner = {c: f for f in data["functionalities"] for c in f["cases"]}
    found = [r for r in recs if r["status"] in ("FAIL", "CONDITIONAL PASS")]
    rows, issues = [], []
    for r in found:
        f = owner.get(r["id"])
        how = coverage(r["id"], f, by_fid)
        if how[0] == "own":
            answer = 'yes: "' + cut(how[1]) + '"' + (" and " + str(how[2] - 1) + " more" if how[2] > 1 else "")
        elif how[0] == "through":
            answer = "yes, through " + how[2] + ", its rule for " + how[3] + ': "' + cut(how[1]) + '"'
        else:
            answer = "none, " + how[1] + ": Catalogue issue to file, below"
            issues.append((r, f, how[1]))
        rows.append([r["id"], r["status"], (f["title"] + " (`" + f["id"] + "`)") if f else "not in functionalities.json", answer])
    module = data.get("module", "").upper()
    lines = ["# Findings, " + module + " " + run + " run", "",
             "Written " + date.today().isoformat() + " from " + Path(path).name + " for a " + entity + " entity. "
             "Each FAIL and CONDITIONAL PASS, its functionality, and the screen requirement that states the check "
             "for it: in its own functionality file, or in the one a row there hands it to.", ""]
    if rows:
        lines += ["| Case | Result | Functionality | Screen requirement stating its check |", "|---|---|---|---|"]
        lines += ["| " + " | ".join(c.replace("|", "/") for c in row) + " |" for row in rows]
    else:
        lines.append("No case is FAIL or CONDITIONAL PASS.")
    lines += ["", "## Catalogue issues to file", ""]
    if not issues:
        lines.append("None: every FAIL and CONDITIONAL PASS has a screen requirement stating its check.")
    else:
        lines += ["Not filed. Paste each into a new issue on the ABDM Developer Portal repository. Write the "
                  "observed line in your own words, with no Aadhaar number, mobile, ABHA number, ABHA address or name.", ""]
        for r, f, why in issues:
            where = (f["title"] + " (" + f["id"] + "), " + f["file"]) if f else "no functionality: the case is not in functionalities.json"
            lines += ["### " + r["id"] + ": no screen requirement in " + (f["id"] if f else "any functionality"), "",
                      "```text",
                      "Title: " + (f["id"] if f else module) + ": no screen requirement states the check for " + r["id"],
                      "Functionality: " + where,
                      "Case: " + r["id"] + ", " + r["status"] + " in a " + run + " run for a " + entity + " entity",
                      "Why: " + why,
                      "Expected Result: " + mask(r["expected"]),
                      "Observed: <one line, in your own words>",
                      "Asked: a screen requirement in the functionality file that states the check for " + r["id"] +
                      ", so the next integrator builds for it before the test.",
                      "```", ""]
    target = Path(out)
    target = target.with_name(target.stem + "-findings.md")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("\n".join(lines).rstrip("\n") + "\n")
    print("findings: " + str(len(found)) + " FAIL or CONDITIONAL PASS, " + str(len(issues)) +
          " without a screen requirement")
    for r, f, _ in issues:
        print("catalogue issue " + r["id"] + " in " + (f["id"] if f else "no functionality"))
    print("wrote " + str(target))


def main():
    p = argparse.ArgumentParser()
    p.add_argument("action", choices=["plan", "check", "build", "tracker", "findings"])
    p.add_argument("workbook")
    p.add_argument("--run", choices=["ft", "ui"], required=True)
    p.add_argument("--entity", choices=["government", "private"])
    p.add_argument("--meta")
    p.add_argument("--evidence", default="abdm-ft/evidence")
    p.add_argument("--out")
    p.add_argument("--manifest")
    p.add_argument("--format", choices=["pdf", "docx", "both"], default="both")
    p.add_argument("--fonts")
    p.add_argument("--draft", action="store_true",
                   help="build even when mandatory cases are not PASS, marked Draft on every page")
    p.add_argument("--functionalities", help="tracker, findings: functionalities.json, else the one beside this script")
    p.add_argument("--plan", default="INTEGRATION-PLAN.md", help="tracker: the plan file the tracker is written into")
    p.add_argument("--state", help="tracker: its state file, else tracker.json beside the workbook")
    p.add_argument("--set", action="append", help="tracker: <functionality id>=building|built|tested|in-scope")
    p.add_argument("--screen", action="append", help="tracker: <case id>=<the screen or route that satisfies it>")
    p.add_argument("--build", help="tracker: the build a functionality was tested on, with --set <id>=tested")
    a = p.parse_args()
    meta = {}
    if a.action == "build" and not (a.meta and a.out):
        sys.exit("build needs --meta and --out")
    if a.meta:
        meta = json.loads(Path(a.meta).read_text())
        bad = forbidden(meta)
        if bad:
            sys.exit("meta.json carries a credential, a field named like one or a value holding one, "
                     "which never goes in a report: " + ", ".join(bad))
        meta["entity_type"] = a.entity or meta.get("entity_type")
    entity = a.entity or meta.get("entity_type")
    if a.action == "tracker":
        run_tracker(a, entity, agreed_cases(meta))
        return
    if entity not in ("government", "private"):
        sys.exit("--entity, or entity_type in meta.json, must be government or private")
    if a.action == "plan":
        run_plan(a.workbook, a.run, entity)
        return
    if a.action == "findings":
        data = load_functionalities(a.functionalities)
        if not data:
            sys.exit("no functionalities.json beside this script: this module's skill is not built per functionality yet")
        if not a.out:
            sys.exit("findings needs --out, the report it sits beside")
        write_findings(a.workbook, a.run, entity, a.out, data)
        return
    recs, groups, blocking = run_check(a.workbook, a.run, entity, agreed_cases(meta))
    if blocking and not (a.action == "build" and a.draft):
        if a.action == "build":
            print("no report: every mandatory case must be PASS first. --draft writes a draft marked as one",
                  file=sys.stderr)
        sys.exit(1)
    if a.action == "build":
        build(a.workbook, a.run, meta, a.evidence, a.out, recs, groups, a.manifest, a.format, a.fonts,
              blocking if blocking else None)
        data = load_functionalities(a.functionalities)
        if data:
            write_findings(a.workbook, a.run, entity, a.out, data, recs)


if __name__ == "__main__":
    main()
