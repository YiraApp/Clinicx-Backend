"""Read NHA's functional test case workbook and write a run's results into it.

  python3 -I abdm_workbook.py cases  WORKBOOK
  python3 -I abdm_workbook.py write  WORKBOOK --run ft|ui --results RESULTS.json
  python3 -I abdm_workbook.py open   WORKBOOK --run ft|ui [--entity government|private]

A result's outcome is passed, conditional-pass, failed, needs-human or not-run.
A conditional-pass needs remarks naming the condition. A result may carry
"steps": [{"step": 1, "outcome": "passed", "note": ""}], appended to Remarks.
Response and Remarks are masked before they are written: Aadhaar, ABHA number,
ABHA address, mobile, JWT and any value after a secret's label.

Needs openpyxl and Pillow (pip install openpyxl pillow); Pillow keeps the
sheet's logos when it is saved.
"""
import argparse, json, re, shutil, sys
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl

# Where each result goes. A ui run fills NHA's own columns, an ft run its own
# four, so neither overwrites the other. The first column a name matches is
# used; a field no name matches gets a column added after the last header.
COLUMNS = {
    "ui": {"status": ["Status(PASS/FAIL)", "Status"], "response": ["Response of Request with Code"],
           "iterations": ["Iterations"], "remarks": ["Remarks", "Remark"]},
    "ft": {"status": ["FT Status"], "response": ["FT Response with Code"],
           "iterations": ["FT Iterations"], "remarks": ["FT Remarks"]},
}
STATUS = {"passed": "PASS", "conditional-pass": "CONDITIONAL PASS", "failed": "FAIL",
          "needs-human": "", "not-run": "NA"}
CLOSED = {"PASS", "CONDITIONAL PASS", "NA"}
# A step of a multi step case, as Remarks shows it.
STEP_STATUS = {"passed": "PASS", "failed": "FAIL", "needs-human": "needs tester", "not-run": "not run"}
# NHA's PHR sheet numbers its rows but gives them no ids. Each tab is one PHR
# milestone, and a case's id is built the way the HMIS ids are: row 3.2 of the
# P1 tab is PHR_ADDR_302, row 5.1.1 of P2 is PHR_APP_50101, a row numbered 9
# alone is 900. The test case page builds the same ids from the same rows.
PHR_TABS = {"ABHA address creation": "ADDR", "PHR App Functionality": "APP",
            "Building HIU Service for PHR": "HIU", "Locker": "LOCKER"}
HEADER_BLOCK = {"Test Executed by:": "executed_by", "Test Execution date:": "date",
                "Application/ System Name:": "application"}


def norm(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


# Masking. This block is the same in abdm_workbook.py and abdm_report.py, so the
# sheet and the report hide the same things in the same way.
# A JWT: three base64url segments, the first a JSON header, so starting eyJ.
JWT = re.compile(r"\beyJ[\w-]+\.[\w-]+\.[\w-]*")
SECRET_WORDS = r"pass(?:word|wd)|secret|token|otp|credentials?|api[ _-]?key|private[ _-]?key"


def labelled(words):
    """A value after a label naming a secret: "client secret: x", "OTP 482913", "otp":"482913",
    "X-token: Bearer x". Only a value is taken, never the { or [ that opens an object or a list."""
    return re.compile(
        r"(\b[\w-]*(?:" + words + r")[\w-]*['\"]?\s*[:=]\s*['\"]?)(?![\[{])((?:Bearer\s+)?[^\s'\",;}]+)"
        r"|(\b(?:" + words + r")\b\s+(?:is\s+)?)((?=[^\s,;]*\d)[^\s,;]{4,})", re.I)


LABELLED = labelled(SECRET_WORDS)
BEARER = re.compile(r"\bBearer\s+(?!\[token\]|\[redacted\])[^\s'\",;}]+")
# A person's profile fields, by key, in JSON or in "key: value" and "key=value" text, in any case.
# A quoted value is taken whole; an unquoted one runs to the end of its clause: a comma, a
# semicolon, a full stop, a line end or a steps list.
PROFILE_KEYS = (r"firstName|middleName|lastName|fullName|name|dob|dayOfBirth|monthOfBirth|yearOfBirth|"
                r"address|pinCode|pincode|photo|profilePhoto|kycPhoto|email")
PROFILE = re.compile(
    r"(['\"]?\b(?i:" + PROFILE_KEYS + r")\b['\"]?\s*[:=]\s*)(?![\[{])"
    r"(\"(?:[^\"\\]|\\.)*\"|'[^']*'|[^\s,;}\]]+(?: (?!Steps:)[^\s,;.}\]]+)*)")
# A date of birth after its label, and any date written dd/mm/yyyy.
DATE = r"\d{1,2}[/-]\d{1,2}[/-]\d{4}|\d{4}-\d{2}-\d{2}"
BIRTH = re.compile(r"(\b(?i:dob|date of birth|birth ?date)\b['\"]?\s*[:=]?\s*['\"]?)(" + DATE + r")(?!\d)")
SLASHED = re.compile(r"(?<![\d/])\d{1,2}/\d{1,2}/\d{4}(?![\d/])")
ABHA_ADDRESS = re.compile(r"(?<![\w.-])([A-Za-z0-9._]+)@(sbx|abdm)\b(?!\.\w)", re.I)
# A digit run is not a number to mask when it is a segment of a UUID.
LEAD, TRAIL = r"(?<!\d)(?<![0-9a-fA-F]-)", r"(?!\d)(?!-[0-9a-fA-F])"
ABHA_NUMBER = re.compile(LEAD + r"\d{2}[ -]?\d{4}[ -]?\d{4}[ -]?(\d{4})" + TRAIL)
AADHAAR = re.compile(LEAD + r"\d{4}[ -]?\d{4}[ -]?(\d{4})" + TRAIL)
MOBILE = re.compile(r"(?<![\d+])(?<![0-9a-fA-F]-)(?:\+91[ -]?|91[ -]?|0)?[6-9](?:[ -]?\d){5}"
                    r"[ -]?(\d)[ -]?(\d)[ -]?(\d)[ -]?(\d)" + TRAIL)


def profile(m):
    quote = m.group(2)[0] if m.group(2)[:1] in ("'", '"') else ""
    return m.group(1) + quote + "[masked]" + quote


def address(m):
    """An ABHA address keeps the last four characters before the @, the rest as x. Four or fewer: all x."""
    name = m.group(1)
    if len(name) <= 4:
        return "x" * len(name) + "@" + m.group(2)
    return "x" * (len(name) - 4) + name[-4:] + "@" + m.group(2)


def mask(text):
    """A JWT to [token], a secret's value and a Bearer credential to [redacted], a profile field's
    value and a date of birth to [masked], a dd/mm/yyyy date to [date], an ABHA address to
    xxxxxx1234@sbx, ABHA numbers to XX-XXXX-XXXX-1234, any other 12 digit run to XXXX XXXX 1234,
    a mobile, with or without 0 or +91, to XXXXXX1234."""
    text = JWT.sub("[token]", str(text))
    text = LABELLED.sub(lambda m: (m.group(1) or m.group(3)) + "[redacted]", text)
    text = BEARER.sub("Bearer [redacted]", text)
    text = BIRTH.sub(lambda m: m.group(1) + "[masked]", text)
    text = PROFILE.sub(profile, text)
    text = SLASHED.sub("[date]", text)
    text = ABHA_ADDRESS.sub(address, text)
    text = ABHA_NUMBER.sub(r"XX-XXXX-XXXX-\1", text)
    text = AADHAAR.sub(r"XXXX XXXX \1", text)
    return MOBILE.sub(r"XXXXXX\1\2\3\4", text)
# End of the masking block.


def steps_text(case, steps):
    """A multi step case's steps as one line: "Steps: 1 PASS; 2 needs tester (what); 3 PASS"."""
    if not isinstance(steps, list):
        sys.exit(case + ": steps must be a list of {\"step\": n, \"outcome\": ..., \"note\": ...}")
    parts = []
    for s in steps:
        if not isinstance(s, dict) or s.get("outcome") not in STEP_STATUS:
            sys.exit(case + ": a step's outcome must be one of " + ", ".join(STEP_STATUS))
        note = mask(norm(s.get("note", ""))).replace(";", ",")
        parts.append(str(s.get("step", len(parts) + 1)) + " " + STEP_STATUS[s["outcome"]] +
                     (" (" + note + ")" if note else ""))
    return "Steps: " + "; ".join(parts) if parts else ""


def steps_agree(outcome, steps):
    """Why a case's outcome does not agree with its steps, or None when it does."""
    seen = [s.get("outcome") for s in steps if isinstance(s, dict)]
    if "failed" in seen and outcome != "failed":
        return "a failed step makes the case failed"
    if outcome == "passed" and any(x != "passed" for x in seen):
        return "a passed case needs every step passed"
    if outcome == "needs-human" and "needs-human" not in seen:
        return "a needs-human case needs a step that is needs-human"
    return None


def load(path):
    # rich_text keeps the bold and coloured runs inside NHA's cells.
    return openpyxl.load_workbook(path, rich_text=True)


def header_row(ws):
    """The header row and, for each name, the first column carrying it."""
    for row in ws.iter_rows(min_row=1, max_row=15):
        names = [norm(c.value) for c in row]
        if "Test Case ID" in names or (norm(ws.title) in PHR_TABS and any(
                re.fullmatch(r"s\.\s*no", n.lower()) for n in names)):
            cols = {}
            for name, cell in zip(names, row):
                if name:
                    cols.setdefault(name, cell.column)
            return row[0].row, cols
    return None, {}


def lost_zero(number, last):
    """A cell typed 2.10 is stored as 2.1, so after 2.9 a row reading 2.1 is 2.10."""
    parts, before = number.split("."), last.split(".")
    if len(parts) == 2 and len(before) == 2 and parts[0] == before[0]:
        if before[1] == str(int(parts[1]) * 10 - 1):
            return f"{parts[0]}.{int(parts[1]) * 10}"
    return number


def phr_cases(ws, top):
    """PHR case ids, from the row numbers, by the same rules as the page."""
    head = [norm(c.value).lower() for c in ws[top]]
    number = next(i for i, n in enumerate(head) if re.fullmatch(r"s\.\s*no", n)) + 1
    content = [i + 1 for i, n in enumerate(head)
               if n.startswith(("test case", "test scenario")) or "steps" in n or "expected" in n]
    down = {}  # a value merged down one column belongs to every row it spans
    for rng in ws.merged_cells.ranges:
        if rng.min_col == rng.max_col and rng.max_row > rng.min_row:
            for r in range(rng.min_row + 1, rng.max_row + 1):
                down[(r, rng.min_col)] = ws.cell(rng.min_row, rng.min_col).value
    found, last = {}, ""
    for r in range(top + 1, ws.max_row + 1):
        # A number merged down marks one case over several rows: only its first row counts.
        num = "" if (r, number) in down else norm(ws.cell(r, number).value)
        filled = any(norm(down.get((r, c), ws.cell(r, c).value)) for c in content)
        if filled and re.fullmatch(r"\d+(\.\d+)*", num):
            num = lost_zero(num, last)
            last = num
            parts = num.split(".")
            digits = parts[0] + ("".join(p.zfill(2) for p in parts[1:]) if len(parts) > 1 else "00")
            found[f"PHR_{PHR_TABS[norm(ws.title)]}_{digits}"] = r
    return found


def cases(ws):
    """Case id to its first row. A repeated id gets #2, #3, as the page does."""
    top, cols = header_row(ws)
    if top is None:
        return {}
    if "Test Case ID" not in cols:
        return phr_cases(ws, top)
    found, col = {}, cols["Test Case ID"]
    for r in range(top + 1, ws.max_row + 1):
        value = norm(ws.cell(r, col).value)
        if not value:
            continue
        base = key = value.replace(" ", "")
        n = 2
        while key in found:
            key, n = base + "#" + str(n), n + 1
        found[key] = r
    return found


def anchor(ws, row, col):
    """A value written inside a merged range belongs in its top left cell."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col)
    return ws.cell(row, col)


def add_column(ws, top, cols, name):
    last = max(cols.values())
    new = last + 1
    head = ws.cell(top, new, name)
    model = ws.cell(top, last)
    for attr in ("font", "fill", "border", "alignment"):
        setattr(head, attr, copy(getattr(model, attr)))
    letter = openpyxl.utils.get_column_letter
    ws.column_dimensions[letter(new)].width = ws.column_dimensions[letter(last)].width or 18
    cols[name] = new
    return new


def column(ws, top, cols, names):
    for name in names:
        if name in cols:
            return cols[name]
    return add_column(ws, top, cols, names[0])


def index(wb):
    rows, twice = {}, []
    for ws in wb.worksheets:
        for case, row in cases(ws).items():
            if case in rows:
                twice.append(case)
            rows.setdefault(case, (ws, row))
    if twice:
        sys.exit("case id on more than one tab: " + ", ".join(twice))
    return rows


def write(path, run, results, meta):
    path = Path(path)
    backup = path.with_name(f"{path.stem}.before-abdm-run{path.suffix}")
    if not backup.exists():
        shutil.copy2(path, backup)
    original = load(backup)
    wb = load(path)
    rows = index(wb)
    where = {}
    for ws in wb.worksheets:
        top, cols = header_row(ws)
        if top is None:
            continue
        if "Test Case ID" not in cols:
            # The PHR sheet has no id column: add one, so the sheet carries the ids it is read by.
            ids = cases(ws)
            col = add_column(ws, top, cols, "Test Case ID")
            for case, row in ids.items():
                ws.cell(row, col).value = case
        where[ws.title] = (top, cols, {k: column(ws, top, cols, names) for k, names in COLUMNS[run].items()})
        # The header block above the table: tester, date, application.
        for line in ws.iter_rows(min_row=1, max_row=top - 1):
            for cell in line:
                text = norm(cell.value)
                for label, key in HEADER_BLOCK.items():
                    if text.startswith(label) and meta.get(key):
                        cell.value = f"{label} {meta[key]}"
    missing = []
    for r in results:
        if r["id"] not in rows:
            missing.append(r["id"])
            continue
        if r["outcome"] not in STATUS:
            sys.exit(f"{r['id']}: outcome must be one of {', '.join(STATUS)}")
        ws, row = rows[r["id"]]
        top, cols, fields = where[ws.title]
        remarks = norm(r.get("remarks", ""))
        if r["outcome"] == "conditional-pass" and not remarks:
            sys.exit(f"{r['id']}: a conditional-pass needs remarks naming the condition")
        if r["outcome"] == "needs-human":
            remarks = f"Needs tester: {remarks}".strip()
        remarks = mask(remarks)
        if r.get("steps"):
            line = steps_text(r["id"], r["steps"])
            wrong = steps_agree(r["outcome"], r["steps"])
            if wrong:
                sys.exit(f"{r['id']}: steps disagree with the outcome {r['outcome']}: {wrong}")
            remarks = " ".join(x for x in (remarks, line) if x)
        values = {"status": STATUS[r["outcome"]], "response": mask(r.get("response") or ""),
                  "iterations": r.get("iterations", ""), "remarks": remarks}
        for key, value in values.items():
            cell = anchor(ws, row, fields[key])
            if norm(original[ws.title][cell.coordinate].value):
                # NHA wrote in this cell. Its text stays; the value goes beside it.
                name = f"{norm(ws.cell(top, fields[key]).value)} (run)"
                cell = anchor(ws, row, cols.get(name) or add_column(ws, top, cols, name))
            cell.value = value
    wb.save(path)
    return missing


# Which cases are mandatory for an entity type: the same reading of the marking
# column and the group's Applicable To as abdm_report.py's plan and check use.
OPTIONAL = ("no", "optional", "non-mandatory", "non mandatory", "not mandatory")


def marking_col(cols):
    return next((col for name, col in sorted(cols.items(), key=lambda x: x[1])
                 if "mandatory" in name.lower() and "optional" in name.lower()), None)



def classify(text, entity):
    """(kind, label, note, pair ids) for one marking and entity type."""
    t = norm(text)
    low = t.lower()
    if not t:
        return "mandatory", "Not marked, treated as mandatory", "no marking: treated as mandatory", None
    if low.startswith("either"):
        ids = [i.replace(" ", "") for i in re.findall(r"\b[A-Z][A-Z0-9]*(?: ?_[A-Z0-9]+)+\b", t)]
        gov, priv = "gov" in low, "priv" in low
        if (entity == "government" and gov) or (entity == "private" and priv) or not (gov or priv):
            return "pair", "Either " + " or ".join(ids), "either rule: one of " + " or ".join(ids) + " must pass", ids
        return "optional", "Optional", "either rule names the other entity type: optional for " + entity, None
    if low in OPTIONAL or low.startswith(("optional", "non-mandatory", "non mandatory", "not mandatory", "no ")):
        return "optional", "Optional", None, None
    m = re.match(r"(mandatory|yes)\b[\s,:]*(.*)$", low)
    if not m:
        return "mandatory", "Unrecognised, treated as mandatory", "unrecognised marking: treated as mandatory", None
    rest = m.group(2)
    if not rest:
        return "mandatory", "Mandatory", None, None
    if rest.startswith("•"):
        # PHR's profile cases: "Mandatory" then a bullet list of the mandatory and optional fields.
        return "mandatory", "Mandatory", "lists mandatory and optional fields: the case is mandatory", None
    conditional = "(" in rest or re.search(r"\bif\b", rest) is not None
    who = re.match(r"for\s+(.*)$", rest)
    if who and m.group(1) == "mandatory":
        head = who.group(1).split("(")[0]
        gov, priv = "gov" in head, "priv" in head
        if (gov or priv) and not ((entity == "government" and gov) or (entity == "private" and priv)):
            return "optional", "Optional", "mandatory for the other entity type only: optional for " + entity, None
        if conditional:
            return "mandatory", "Mandatory (conditional)", "conditional: treated as mandatory", None
        return "mandatory", "Mandatory", "mandatory for " + entity, None
    if conditional:
        return "mandatory", "Mandatory (conditional)", "conditional: treated as mandatory", None
    return "mandatory", "Unrecognised, treated as mandatory", "unrecognised marking: treated as mandatory", None



def group_applies(ws, top, cols):
    """Each case row's group Applicable To, carried down from the row that heads the group."""
    acol, icol = cols.get("Applicable To"), cols.get("Test Case ID")
    if not acol or not icol:
        return {}
    out, current = {}, ""
    for r in range(top + 1, ws.max_row + 1):
        applies = norm(anchor(ws, r, acol).value)
        if applies:
            current = applies
        if norm(ws.cell(r, icol).value):
            out[r] = current
    return out



def restricted_to(applies):
    """The one entity type a group is for, when it names only one: government or private."""
    low = applies.lower()
    gov, priv = "gov" in low, "priv" in low
    if "all" in low or "both" in low or gov == priv:
        return None
    return "government" if gov else "private"



def status_of(ws, row, col):
    return norm(anchor(ws, row, col).value)


def still_open(path, run, entity=None):
    """Cases not yet PASS, CONDITIONAL PASS or NA in this run type's status column. With an entity
    type, also an NA on a case mandatory for it, as the report's gate blocks one."""
    path = Path(path)
    wb = load(path)
    backup = path.with_name(f"{path.stem}.before-abdm-run{path.suffix}")
    original = load(backup) if backup.exists() else None
    found, applies_by = [], {}
    for case, (ws, row) in index(wb).items():
        top, cols = header_row(ws)
        col = next((cols[n] for n in COLUMNS[run]["status"] if n in cols), None)
        status = ""
        if col is not None:
            cell = anchor(ws, row, col)
            run_col = cols.get(f"{norm(ws.cell(top, col).value)} (run)")
            if run_col and original is not None and norm(original[ws.title][cell.coordinate].value):
                # NHA wrote in this cell, so the run's status is only in the "(run)" column.
                status = status_of(ws, row, run_col)
            else:
                status = norm(cell.value)
        kind, pair = "mandatory", None
        if entity:
            if ws.title not in applies_by:
                applies_by[ws.title] = group_applies(ws, top, cols)
            mcol = marking_col(cols)
            kind, _, _, pair = classify(norm(anchor(ws, row, mcol).value) if mcol else "", entity)
            only = restricted_to(applies_by[ws.title].get(row, ""))
            if only and only != entity:
                kind, pair = "optional", None
        found.append((case, status.upper(), kind, pair))
    passing = {case for case, status, _, _ in found if status in ("PASS", "CONDITIONAL PASS")}
    paired = {i for _, _, kind, pair in found if kind == "pair" for i in pair}
    for case, status, kind, pair in found:
        if status not in CLOSED:
            print(case)
        elif entity and status == "NA":
            if case in paired:
                ids = next((p for _, _, k, p in found if k == "pair" and case in p), [case])
                if not any(i in passing for i in ids):
                    print(case + " (NA on a mandatory case, either " + " or ".join(ids) + ")")
            elif kind == "mandatory":
                print(case + " (NA on a mandatory case)")


def main():
    p = argparse.ArgumentParser()
    p.add_argument("action", choices=["cases", "write", "open"])
    p.add_argument("workbook")
    p.add_argument("--run", choices=["ft", "ui"])
    p.add_argument("--results")
    p.add_argument("--executed-by", default="")
    p.add_argument("--application", default="")
    p.add_argument("--entity", choices=["government", "private"],
                   help="open: also list an NA on a case mandatory for this entity type")
    a = p.parse_args()
    if a.action == "cases":
        wb = load(a.workbook)
        print(json.dumps({ws.title: list(cases(ws)) for ws in wb.worksheets}, indent=1))
    elif a.action == "open":
        still_open(a.workbook, a.run or sys.exit("--run is needed"), a.entity)
    else:
        if not (a.run and a.results):
            sys.exit("write needs --run and --results")
        meta = {"executed_by": a.executed_by, "application": a.application,
                "date": date.today().isoformat()}
        missing = write(a.workbook, a.run, json.loads(Path(a.results).read_text()), meta)
        for case in missing:
            print(f"not in the workbook: {case}", file=sys.stderr)
        sys.exit(1 if missing else 0)


if __name__ == "__main__":
    main()
